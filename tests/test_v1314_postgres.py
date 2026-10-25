"""Destructive integration tests: explicit loopback *_test database ONLY."""
import io
import json
import os
import shutil
import sqlite3
from datetime import datetime, timedelta
from pathlib import Path
from urllib.parse import urlsplit

import openpyxl
import pytest


def test_turso_configuration_is_fail_closed(monkeypatch):
    from psx_app.database.turso import credentials, enabled
    monkeypatch.setenv('PSX_DATABASE_BACKEND', 'turso')
    monkeypatch.delenv('TURSO_AUTH_TOKEN', raising=False)
    monkeypatch.setenv('TURSO_DATABASE_URL', 'libsql://example.turso.io')
    assert enabled()
    with pytest.raises(RuntimeError, match='AUTH_TOKEN'):
        credentials()
    monkeypatch.setenv('TURSO_AUTH_TOKEN', 'synthetic-test-only')
    monkeypatch.setenv('TURSO_DATABASE_URL', 'http://example.com')
    with pytest.raises(RuntimeError, match='valid'):
        credentials()


@pytest.fixture
def remote(monkeypatch,tmp_path):
    url=os.getenv('PSX_TEST_POSTGRES_URL')
    if not url: pytest.skip('Explicit disposable loopback PostgreSQL test server not supplied')
    p=urlsplit(url)
    assert p.hostname in ('127.0.0.1','localhost') and p.path.strip('/').endswith('_test'), 'Only loopback *_test databases allowed'
    import psycopg
    from psx_app.database import postgres
    with psycopg.connect(url,autocommit=True) as raw:
        raw.execute('DROP SCHEMA public CASCADE');raw.execute('CREATE SCHEMA public')
    monkeypatch.setenv('PSX_DATABASE_BACKEND','postgres');monkeypatch.setenv('DATABASE_URL',url);monkeypatch.setenv('PSX_PG_LOOPBACK_TEST','1')
    monkeypatch.setattr(postgres,'_READY',False);postgres.ensure_ready()
    import psx_app.main as main
    main._API_CACHE.clear();main._READ_FRESHNESS.clear()
    from psx_app.database.availability import readiness
    readiness.failure=None;readiness.retry_at=0
    yield tmp_path


def workbook(rows=20, close=100):
    book = openpyxl.Workbook()
    book.active.append(['DATE','SYMBOL','VOLUME','PR-CLOSE','HIGH','LOW','OPEN','CLOSE','AVG'])
    for i in range(rows):
        book.active.append([datetime(2025,1,1)+timedelta(days=i),'ATRL',1000,99,103,98,99,close,100])
    data = io.BytesIO()
    book.save(data)
    return data.getvalue()


def test_remote_commit_rollback_mapping_and_revision(remote):
    from psx_app.database.schema import get_connection
    from psx_app.database.turso import revision
    before = revision()
    with get_connection() as c:
        row = c.execute('SELECT 12 AS n, ? AS label', ('test',)).fetchone()
        assert row[0] == row['n'] == 12 and dict(row)['label'] == 'test'
        c.execute("INSERT INTO audit_events(event_type,details_json,created_at) VALUES ('TEST','{}','2025')")
    assert revision() != before
    before = revision()
    with pytest.raises(ValueError):
        with get_connection() as c:
            c.execute("INSERT INTO audit_events(event_type,details_json,created_at) VALUES ('ROLLBACK','{}','2025')")
            raise ValueError('test rollback')
    with get_connection() as c:
        assert c.execute("SELECT COUNT(*) FROM audit_events WHERE event_type='ROLLBACK'").fetchone()[0] == 0
    assert revision() == before


def test_remote_bulk_insert_and_cache_does_not_change_revision(remote):
    from psx_app.database.schema import get_connection
    from psx_app.database.turso import revision
    before = revision()
    with get_connection() as c:
        c.executemany('INSERT INTO indicator_cache(symbol,trade_date,payload_json,updated_at) VALUES (?,?,?,?)',
                      [('T'+str(i),'2025-01-01','{}','2025') for i in range(2000)])
    with get_connection() as c:
        assert c.execute('SELECT COUNT(*) FROM indicator_cache').fetchone()[0] == 2000
    assert revision() == before


def test_remote_stage_commit_on_different_worker_and_idempotency(remote):
    from psx_app.services.import_orchestrator import ImportOrchestrator
    from psx_app.database.schema import get_connection
    first = ImportOrchestrator(data_root=remote/'worker-a')
    stage = first.stage_uploaded_files([('2025.xlsx', workbook())])
    shutil.rmtree(remote/'worker-a', ignore_errors=True)
    second = ImportOrchestrator(data_root=remote/'worker-b')
    committed = second.commit_staged_batch(stage['batch_uuid'])
    assert committed['new_records'] == 20
    assert second.commit_staged_batch(stage['batch_uuid'])['status'] == 'ALREADY_COMMITTED'
    with get_connection() as c:
        assert c.execute('SELECT COUNT(*) FROM equity_ohlcv').fetchone()[0] == 20
        assert c.execute('SELECT COUNT(*) FROM psx_blob_parts').fetchone()[0] == 0
        assert c.execute('SELECT archive_path FROM source_files').fetchone()[0].startswith('remote:metadata-only/')
    again = first.stage_uploaded_files([('2025.xlsx', workbook())])
    assert again['duplicates_ignored'] == 20
    assert first.commit_staged_batch(again['batch_uuid'])['new_records'] == 0
    revised = first.stage_uploaded_files([('2025.xlsx', workbook(close=101))])
    assert revised['revisions_count'] == 20
    result = first.commit_staged_batch(revised['batch_uuid'])
    assert result['revisions_count'] == 20


def test_remote_blob_cross_worker_parts_and_expiry(remote):
    from psx_app.services.remote_staging import put, get, delete, sweep
    from psx_app.database.schema import get_connection
    data = bytes(range(256)) * 6000
    put('chunk:test:0', data)
    assert get('chunk:test:0') == data
    with get_connection() as c:
        c.execute('UPDATE psx_blob_parts SET updated_at=0 WHERE scope=?', ('chunk:test:0',))
    sweep()
    with pytest.raises(FileNotFoundError):
        get('chunk:test:0')


def test_remote_http_endpoints_and_no_snapshot_fallback(remote, monkeypatch):
    import psx_app.main as main
    from fastapi.testclient import TestClient
    from psx_app.services.backup_service import BackupService
    from psx_app.services.cloud_sync_service import LegacyCloudSyncService
    monkeypatch.setattr(main, 'DATA_ROOT', remote)
    monkeypatch.setattr(LegacyCloudSyncService, '_gh_request', lambda *a, **kw: pytest.fail('GitHub must not be contacted'))
    main._API_CACHE.clear()
    main._READ_FRESHNESS.clear()
    with TestClient(main.app) as client:
        for path in ('/api/health','/api/overview','/api/settings','/api/data-dates','/api/accounts/91404','/api/futures/comparison'):
            response = client.get(path)
            assert response.status_code == 200, (path, response.text[:300])
        status = client.get('/api/settings').json()
        assert status['cloud_sync']['backend'] == 'postgres'
        assert status['db_integrity']['check_type'] == 'connectivity'
        assert client.get('/api/cloud-sync/download-snapshot').status_code == 409
        assert client.post('/api/backups/create', json={}).status_code == 409
        assert client.post('/api/admin/date-normalization/plan').status_code == 409
        result = client.post('/api/settings', json={'updates': {'futures_default_days': 20}})
        assert result.status_code == 200 and result.json()['cloud_sync']['saved']
        assert client.post('/api/cloud-sync/reload', json={'confirm':True}).status_code == 200
    with pytest.raises(ValueError, match='NEW remote'):
        BackupService().restore_backup('anything.zip')
    assert not list(remote.rglob('app.db'))


def test_remote_account_rollups_are_atomic_and_separate(remote, monkeypatch):
    from psx_app.services.import_orchestrator import ImportOrchestrator
    from psx_app.services.account_service import AccountService
    from psx_app.database.schema import get_connection
    header = 'Date,Time,Terminal,Market,Symbol,Transaction Type,Order ID,Execution ID,Account,Status,Quantity,Remaining Quantity,Order Quantity,Trade Price,Trade Value\n'
    data = (header + '16/10/2025,10:00:01,THM00601,REG,OGDC,Buy,O1,E1,91404,FILLED,1000,0,1000,200,200000\n' +
            '16/10/2025,11:00:01,THM00601,REG,OGDC,Sell,O2,E2,91404,FILLED,1000,0,1000,202,202000\n').encode()
    orch = ImportOrchestrator(data_root=remote)
    stage = orch.stage_uploaded_files([('TradeLog.csv', data)])
    assert stage['rows_accepted'] == 2
    original = AccountService.recalculate_account_date_from_executions
    def fail(*a, **kw):
        raise ValueError('synthetic account rollup failure')
    monkeypatch.setattr(AccountService, 'recalculate_account_date_from_executions', fail)
    with pytest.raises(ValueError, match='rollup failure'):
        orch.commit_staged_batch(stage['batch_uuid'])
    with get_connection() as c:
        assert c.execute('SELECT COUNT(*) FROM trade_executions').fetchone()[0] == 0
        assert c.execute('SELECT COUNT(*) FROM account_daily_symbol_summary').fetchone()[0] == 0
        assert c.execute('SELECT status FROM import_batches').fetchone()[0] == 'FAILED'
    monkeypatch.setattr(AccountService, 'recalculate_account_date_from_executions', original)
    assert orch.commit_staged_batch(stage['batch_uuid'])['new_records'] == 2
    with get_connection() as c:
        row = c.execute('SELECT account_id,net_pnl FROM account_daily_symbol_summary').fetchone()
        assert row['account_id'] == '91404'
        assert row['net_pnl'] == pytest.approx(2000 - 202000 * 0.000325 * 1.13)
        assert c.execute("SELECT COUNT(*) FROM account_daily_symbol_summary WHERE account_id='91269'").fetchone()[0] == 0


def test_remote_chunked_upload_and_commit(remote, monkeypatch):
    import psx_app.main as main
    from fastapi.testclient import TestClient
    from psx_app.services.import_orchestrator import ImportOrchestrator
    monkeypatch.setattr(main, 'DATA_ROOT', remote/'worker')
    main._API_CACHE.clear(); main._READ_FRESHNESS.clear()
    data = workbook(2)
    upload_id = 'remote-chunk-test'
    pieces = [data[:len(data)//2], data[len(data)//2:]]
    with TestClient(main.app) as client:
        for i, part in enumerate(pieces):
            response = client.post(f'/api/imports/upload-chunk?upload_id={upload_id}&name=2025.xlsx&idx={i}&total=2', content=part)
            assert response.status_code == 200 and response.json()['storage'] == 'remote'
        shutil.rmtree(remote/'worker', ignore_errors=True)
        response = client.post('/api/imports/stage-chunked', json={'uploads':[{'upload_id':upload_id,'name':'2025.xlsx','total_chunks':2}]})
        assert response.status_code == 200, response.text
        batch = response.json()['batch_uuid']
        history = client.get('/api/imports/history').json()
        assert any(b['batch_uuid']==batch and b['staged_resumable'] for b in history['batches'])
        response = client.post('/api/imports/commit', json={'batch_uuid':batch})
        assert response.status_code == 200, response.text
        assert response.json()['cloud_sync']['saved']


def test_remote_freshness_invalidates_cached_dates(remote):
    import psx_app.main as main
    from fastapi.testclient import TestClient
    from psx_app.database.schema import get_connection
    main._API_CACHE.clear(); main._READ_FRESHNESS.clear()
    with TestClient(main.app) as client:
        assert client.get('/api/data-dates').json()['cash'] is None
        with get_connection() as c:
            c.execute("INSERT INTO equity_ohlcv(trade_date,symbol,close) VALUES ('2025-10-22','ATRL',100)")
        result = client.get('/api/system/freshness?force=true')
        assert result.status_code == 200 and result.json()['changed']
        assert client.get('/api/data-dates').json()['cash'] == '2025-10-22'


def test_remote_v139_reference_forecast_learning_tables_and_revision(remote):
    from psx_app.database.schema import get_connection
    from psx_app.database.turso import revision
    from psx_app.services.sector_reference import SectorReference
    from psx_app.services.forecast_ledger import ForecastLedger
    from psx_app.services.learning_research import train_symbol
    import math
    with get_connection() as c:
        c.executemany('INSERT INTO equity_ohlcv(symbol,trade_date,open,high,low,close,pr_close,volume) VALUES (?,?,?,?,?,?,?,?)',[
            ('ATRL',(datetime(2024,1,1)+timedelta(days=i)).date().isoformat(),100+i*.04,102+i*.04,99+i*.04,100+i*.04,99.96+i*.04,10000) for i in range(340)])
    before=revision()
    ledger=ForecastLedger();ledger.record('ATRL');ledger.record('ATRL')
    assert len(ledger.list()['items'])==2
    assert ledger.reconcile()['evaluated']==2
    model=train_symbol('ATRL');assert model['status']=='RESEARCH_ONLY'
    # Forecast/research metadata must not invalidate the canonical price revision on each view.
    assert revision()==before
    raw=io.BytesIO();wb=openpyxl.Workbook();ws=wb.active;ws.title='COMP_LIST';ws.append(['SYMBOL','COMPANY NAME','SECTOR']);ws.append(['ATRL','Attock Refinery','REFINERY']);wb.save(raw)
    service=SectorReference();p=service.preview(raw.getvalue());assert revision()==before
    service.commit(p['token']);assert revision()!=before
    assert service.directory()['ATRL']['sector']=='REFINERY'


def reference_workbook_1310(rows):
    w=openpyxl.Workbook();s=w.active;s.title='COMP_LIST'
    s.append(['SYMBOL','COMPANY NAME','SECTOR'])
    for row in rows:s.append(row)
    b=io.BytesIO();w.save(b);return b.getvalue()


def test_remote_reference_normal_upload_across_workers(remote):
    from psx_app.services.import_orchestrator import ImportOrchestrator
    from psx_app.database.schema import get_connection
    a=ImportOrchestrator(data_root=remote/'worker-a')
    raw=reference_workbook_1310([['TRG','TRG Pakistan','TECHNOLOGY']])
    staged=a.stage_uploaded_files([('SECTOR_WISE.xlsx',raw)])
    assert staged['rows_accepted']==1
    shutil.rmtree(remote/'worker-a'/'data'/'staging')
    b=ImportOrchestrator(data_root=remote/'worker-b')
    result=b.commit_staged_batch(staged['batch_uuid'])
    assert result['status']=='COMMITTED' and result['new_records']==1
    assert b.commit_staged_batch(staged['batch_uuid'])['status']=='ALREADY_COMMITTED'
    m=b.stage_uploaded_files([('renamed.xlsx',reference_workbook_1310([['TRG','Changed Company','TECHNOLOGY']]))])
    assert b.commit_staged_batch(m['batch_uuid'])['revisions_count']==1
    with get_connection() as c:
        assert c.execute("SELECT company FROM company_reference WHERE symbol='TRG'").fetchone()[0]=='Changed Company'
        assert c.execute('SELECT COUNT(*) FROM equity_ohlcv').fetchone()[0]==0
        assert c.execute("SELECT COUNT(*) FROM source_revisions WHERE table_name='company_reference'").fetchone()[0]==1


def test_remote_reference_atomic_rollback(remote):
    from psx_app.services.import_orchestrator import ImportOrchestrator
    from psx_app.database.schema import get_connection
    with get_connection() as c:
        c.execute("ALTER TABLE company_reference ADD CONSTRAINT reject_reference CHECK(symbol<>'FAIL')")
        c.commit()
    service=ImportOrchestrator(data_root=remote)
    m=service.stage_uploaded_files([('sector.xlsx',reference_workbook_1310([['TRG','TRG Pakistan','TECHNOLOGY'],['FAIL','Synthetic Failure','TEST']]))])
    with pytest.raises(Exception,match='transaction was not confirmed|Database connection unavailable'):
        service.commit_staged_batch(m['batch_uuid'])
    with get_connection() as c:
        assert c.execute('SELECT COUNT(*) FROM company_reference').fetchone()[0]==0
        assert c.execute('SELECT COUNT(*) FROM source_files').fetchone()[0]==0


def test_remote_complete_bundled_reference_workbook_is_batched(remote,monkeypatch):
    from psx_app.services.sector_reference import BUNDLED
    from psx_app.services.import_orchestrator import ImportOrchestrator
    from psx_app.database import postgres
    from psx_app.database.schema import get_connection
    book=openpyxl.Workbook();book.remove(book.active)
    lists=json.loads(BUNDLED.read_text());total=sum(len(rows) for rows in lists.values())
    for name,rows in lists.items():
        ws=book.create_sheet(name);ws.append(['SYMBOL','COMPANY NAME','SECTOR'])
        for r in rows:ws.append([r['symbol'],r['company'],r['sector']])
    stream=io.BytesIO();book.save(stream)
    original=postgres.Connection.execute;writes=[]
    def counting(self,sql,parameters=()):
        if sql.startswith('INSERT OR REPLACE INTO company_reference'):writes.append(1)
        return original(self,sql,parameters)
    monkeypatch.setattr(postgres.Connection,'execute',counting)
    service=ImportOrchestrator(data_root=remote)
    staged=service.stage_uploaded_files([('SECTOR_WISE.xlsx',stream.getvalue())])
    assert staged['rows_accepted']==total==903
    result=service.commit_staged_batch(staged['batch_uuid'])
    assert result['new_records']==903
    assert 1 <= len(writes) <= 4
    with get_connection() as c:assert c.execute('SELECT COUNT(*) FROM company_reference').fetchone()[0]==903


def test_remote_market_volume_derived_sum_matches_cash_universe(remote):
    from psx_app.database.schema import get_connection
    from psx_app.services.market_explorer import MarketExplorer
    with get_connection() as c:
        c.executemany('INSERT INTO equity_ohlcv(trade_date,symbol,close,volume) VALUES (?,?,?,?)', [('2026-10-23','ATRL',100,1000),('2026-10-23','PRL',50,2000)])
        c.execute('INSERT INTO index_ohlcv(trade_date,close) VALUES (?,?)',('2026-10-23',100000))
    r=MarketExplorer().symbol('ATRL',1,'2026-10-23')
    assert r['index_rows'][0]['volume']==3000
    assert r['index_rows'][0]['volume_symbols']==2
    assert MarketExplorer().volume('2026-10-23')['total']==3000


def test_remote_cold_start_defers_database_to_worker_and_serves_shell(remote, monkeypatch):
    from fastapi.testclient import TestClient
    from psx_app import main
    from psx_app.database import postgres
    from psx_app.database.availability import Readiness
    monkeypatch.setattr(postgres, '_READY', False)
    monkeypatch.setattr(main, 'readiness', Readiness())
    with TestClient(main.app) as client:
        assert postgres._READY is False
        assert client.get('/').status_code == 200
        assert postgres._READY is False
        health = client.get('/api/health')
        assert health.status_code == 200 and health.json()['db_integrity']['ok']
        assert postgres._READY is True
        assert client.get('/api/v139/market/meta').status_code == 200


def add_prices(symbol='ATRL',count=340):
    from psx_app.database.schema import get_connection
    with get_connection() as c:
        c.executemany('INSERT INTO equity_ohlcv(symbol,trade_date,open,high,low,close,pr_close,avg_price,volume) VALUES (?,?,?,?,?,?,?,?,?)',[
            (symbol,(datetime(2024,1,1)+timedelta(days=i)).date().isoformat(),100+i*.04,102+i*.04,99+i*.04,100+i*.04,99.96+i*.04,100+i*.04,3_000_000_000+i) for i in range(count)])


def test_pg_big_volumes_and_shared_full_analysis(remote,monkeypatch):
    from psx_app.services.technical_service import TechnicalService
    from psx_app.services.market_explorer import MarketExplorer
    from psx_app.database.turso import revision
    add_prices();calls=[]
    original=TechnicalService.calculate_series_for_bars
    def count(self,*a,**kw):calls.append(1);return original(self,*a,**kw)
    monkeypatch.setattr(TechnicalService,'calculate_series_for_bars',count)
    before=revision();service=TechnicalService()
    a=service.get_ticker_analysis('ATRL',1)
    b=TechnicalService().get_ticker_analysis('ATRL',120)
    c=TechnicalService().get_ticker_analysis('ATRL',0)
    assert len(calls)==1 and len(a['series'])==1 and len(b['series'])==120 and len(c['series'])==340
    assert a['latest']['close']==b['latest']['close']
    assert revision()==before
    r=MarketExplorer().volume();assert r['total']>2**31
    json.dumps(r) # no Decimal in JSON result


def test_pg_changed_symbol_only_invalidates_indicators(remote):
    from psx_app.services.technical_service import TechnicalService
    from psx_app.services.import_orchestrator import ImportOrchestrator
    from psx_app.database.schema import get_connection
    add_prices();add_prices('PRL')
    TechnicalService().get_universe_summary(scope='all')
    with get_connection() as c:
        assert c.execute('SELECT COUNT(DISTINCT symbol) FROM indicator_cache').fetchone()[0]==2
    orch=ImportOrchestrator(data_root=remote)
    stage=orch.stage_uploaded_files([('2025.xlsx',workbook(2,close=120))]);orch.commit_staged_batch(stage['batch_uuid'])
    with get_connection() as c:
        assert {r[0] for r in c.execute('SELECT DISTINCT symbol FROM indicator_cache')}=={'PRL'}
    result=TechnicalService().get_ticker_analysis('ATRL',1)
    assert result['latest']['close']==120


def test_pg_all_module_routes_and_exports(remote,monkeypatch):
    from fastapi.testclient import TestClient
    import psx_app.main as main
    from psx_app.database.schema import get_connection
    add_prices()
    with get_connection() as c:
        c.execute("INSERT INTO futures_ohlcv(trade_date,contract_symbol,base_symbol,contract_month,volume,pr_close,open,high,low,close) VALUES ('2024-12-05','ATRL-DEC','ATRL','DEC',500,113,113,115,112,114)")
        c.execute("INSERT INTO index_ohlcv(trade_date,pr_close,open,high,low,close) VALUES ('2024-12-05',100000,100000,100500,99900,100300)")
        c.execute("INSERT INTO top5_volume_records(trade_date,ticker,rank,rate,volume,trades,imported_at) VALUES ('2024-12-05','ATRL',1,114,1000,3,'2024')")
    monkeypatch.setattr(main,'DATA_ROOT',remote)
    paths=['/api/overview','/api/top5/register','/api/top5/trend/ATRL','/api/accounts/91404','/api/accounts/91269','/api/futures/metadata','/api/futures/comparison?apply_rolling=false',
    '/api/futures/symbol-summary','/api/futures/date-summary','/api/technical/universe','/api/technical/ticker/ATRL',
    '/api/intelligence/filters','/api/intelligence/trading-ideas?query=ATRL&scope_or_sector=ALL_MARKET', '/api/intelligence/automated-analyst/ATRL',
    '/api/v139/market/meta','/api/v139/market/volume','/api/v139/market/symbol/ATRL','/api/v139/coverage?search=ATRL','/api/v139/contract/ATRL','/api/v139/learning/ATRL','/api/v139/forecasts','/api/reconciliation','/api/system/status']
    with TestClient(main.app) as client:
        for path in paths:
            res=client.get(path);assert res.status_code==200,(path,res.text[:200])
        for target in ('top5_xlsx','account_91404_xlsx','account_91269_xlsx','futures_xlsx','technical_xlsx','csv_equity'):
            res=client.get('/api/exports/generate/'+target)
            assert res.status_code==200,(target,res.text[:200])


def test_pg_snapshot_migration_preserves_ids_and_refuses_nonempty(remote,tmp_path):
    from psx_app.database.schema import DDL_STATEMENTS,get_connection
    from tools.import_sqlite_to_postgres import migrate
    source=tmp_path/'source.db'
    with sqlite3.connect(source) as c:
        c.executescript(DDL_STATEMENTS)
        c.execute("INSERT INTO equity_ohlcv(id,trade_date,symbol,volume,close) VALUES (42,'2025-01-01','ATRL',3000000000,123.45)")
        c.execute("INSERT INTO accounts VALUES ('91404','THM00601','USB account','PKR',1,'2025')")
        c.execute("INSERT INTO app_settings(key,value_json,updated_at) VALUES ('currency','\"PKR\"','2025')")
    assert migrate(source)['status']=='DRY_RUN'
    result=migrate(source,True);assert result['copied']['equity_ohlcv']==1
    with get_connection() as c:
        r=c.execute('SELECT id,close,volume FROM equity_ohlcv').fetchone();assert tuple(r)==(42,123.45,3000000000)
        c.execute("INSERT INTO equity_ohlcv(trade_date,symbol,close) VALUES ('2025-01-02','ATRL',124)")
        assert c.execute("SELECT MAX(id) FROM equity_ohlcv").fetchone()[0]==43
    with pytest.raises(ValueError,match='already contains'):migrate(source,True)


def test_overview_shared_cache_invalidates_with_canonical_revision(remote, monkeypatch):
    from psx_app.database.schema import get_connection
    from psx_app.database.repository import Repository
    from psx_app.database.postgres import Connection
    counts={'overview_queries':0}
    original=Connection.execute
    def tracked(self,sql,parameters=()):
        if 'SELECT COUNT(*) FROM equity_ohlcv' in sql:counts['overview_queries']+=1
        return original(self,sql,parameters)
    monkeypatch.setattr(Connection,'execute',tracked)
    assert Repository().get_system_overview()['equity_ohlcv']['records']==0
    assert Repository().get_system_overview()['equity_ohlcv']['records']==0
    assert counts['overview_queries']==1
    with get_connection() as c:
        c.execute("INSERT INTO equity_ohlcv(trade_date,symbol,close,volume) VALUES ('2026-01-01','ATRL',100,10)")
    assert Repository().get_system_overview()['equity_ohlcv']['records']==1
    assert counts['overview_queries']==2


def test_cold_universe_history_is_bounded_batch_not_one_query_per_symbol(remote,monkeypatch):
    from psx_app.database.schema import get_connection
    from psx_app.database.postgres import Connection
    from psx_app.services.technical_service import TechnicalService
    symbols=['ATRL','PRL','TRG']
    with get_connection() as c:
        for sym in symbols:
            c.executemany('INSERT INTO equity_ohlcv(trade_date,symbol,close,open,high,low,pr_close,volume) VALUES (?,?,?,?,?,?,?,?)',
                [(f'2026-01-{i:02d}',sym,100+i,100,103+i,99,99+i,1000) for i in range(1,29)])
    queries=[];original=Connection.execute
    def tracked(self,sql,parameters=()):
        if 'SELECT symbol,trade_date,open,high,low,close,pr_close,avg_price,volume FROM equity_ohlcv WHERE symbol IN' in sql:queries.append(sql)
        return original(self,sql,parameters)
    monkeypatch.setattr(Connection,'execute',tracked)
    result=TechnicalService().get_universe_summary(scope='all')
    assert result['count']==3 and len(queries)==1


def test_private_staged_snapshot_checksum_failure_preserves_chunks(remote):
    from psx_app.services.remote_staging import put,get
    from tools.migrate_staged_snapshot import migrate_staged
    put('chunk:private-checksum-test:0',b'not-a-snapshot')
    with pytest.raises(ValueError,match='checksum mismatch'):
        migrate_staged('private-checksum-test',1,'0'*64)
    assert get('chunk:private-checksum-test:0')==b'not-a-snapshot'


def test_private_staged_snapshot_success_copies_and_removes_only_own_chunks(remote,tmp_path):
    import gzip,hashlib
    from psx_app.database.schema import DDL_STATEMENTS,get_connection
    from psx_app.services.remote_staging import put,get
    from tools.migrate_staged_snapshot import migrate_staged
    source=tmp_path/'private.db'
    with sqlite3.connect(source) as c:
        c.executescript(DDL_STATEMENTS)
        c.execute("INSERT INTO equity_ohlcv(id,trade_date,symbol,volume,close) VALUES (52,'2025-01-01','ATRL',3000000000,123.45)")
    c.close()  # Flush any WAL before reading the standalone snapshot bytes.
    payload=gzip.compress(source.read_bytes())
    put('chunk:private-success-test:0',payload)
    put('chunk:unrelated-upload:0',b'retain')
    result=migrate_staged('private-success-test',1,hashlib.sha256(payload).hexdigest())
    assert result['status']=='COMMITTED' and result['temporary_chunks_removed']
    with get_connection() as c:
        assert tuple(c.execute('SELECT id,close,volume FROM equity_ohlcv').fetchone())==(52,123.45,3000000000)
    assert get('chunk:unrelated-upload:0')==b'retain'
    with pytest.raises(Exception):get('chunk:private-success-test:0')


def test_index_import_is_batched_and_keeps_first_value_and_existing_dates(remote,monkeypatch):
    from psx_app.services.import_orchestrator import ImportOrchestrator
    from psx_app.database import postgres
    from psx_app.database.schema import get_connection
    with get_connection() as c:
        c.execute("INSERT INTO index_ohlcv(trade_date,close) VALUES ('2025-01-01',88)")
    raw=workbook(100)
    orch=ImportOrchestrator(data_root=remote)
    stage=orch.stage_uploaded_files([('INDEX.xlsx',raw),('INDEX-copy.xlsx',raw)])
    assert stage['batch_uuid']
    calls=[];original=postgres.Connection.execute
    def counted(self,sql,*args,**kwargs):
        calls.append(sql)
        return original(self,sql,*args,**kwargs)
    monkeypatch.setattr(postgres.Connection,'execute',counted)
    result=orch.commit_staged_batch(stage['batch_uuid'])
    selects=[q for q in calls if q.startswith('SELECT * FROM index_ohlcv WHERE trade_date IN')]
    inserts=[q for q in calls if 'INSERT INTO index_ohlcv' in q]
    assert len(selects)<=12 and len(inserts)==1
    assert result['new_records']==99 and result['duplicates_ignored']>=101
    with get_connection() as c:
        assert c.execute('SELECT COUNT(*) FROM index_ohlcv').fetchone()[0]==100
        assert c.execute("SELECT close FROM index_ohlcv WHERE trade_date='2025-01-01'").fetchone()[0]==88
    assert orch.commit_staged_batch(stage['batch_uuid'])['status']=='ALREADY_COMMITTED'


def test_many_symbol_import_invalidates_in_one_bounded_delete(remote,monkeypatch):
    from psx_app.database import postgres
    from psx_app.database.schema import get_connection
    from psx_app.services.import_orchestrator import ImportOrchestrator
    book=openpyxl.Workbook();book.active.append(['DATE','SYMBOL','VOLUME','PR-CLOSE','HIGH','LOW','OPEN','CLOSE','AVG'])
    symbols=['TEST'+str(i) for i in range(120)]
    for symbol in symbols:book.active.append([datetime(2025,1,2),symbol,1000,99,103,98,99,100,100])
    raw=io.BytesIO();book.save(raw)
    with get_connection() as c:
        c.executemany('INSERT INTO indicator_cache(symbol,trade_date,payload_json,updated_at) VALUES (?,?,?,?)',[(s,'2025-01-01','{}','2025') for s in symbols+['UNTOUCHED']])
    orch=ImportOrchestrator(data_root=remote);stage=orch.stage_uploaded_files([('2025-cash.xlsx',raw.getvalue())])
    queries=[];original=postgres.Connection.execute
    def counted(self,sql,*args,**kwargs):
        if sql.startswith('DELETE FROM indicator_cache'):queries.append(sql)
        return original(self,sql,*args,**kwargs)
    monkeypatch.setattr(postgres.Connection,'execute',counted)
    result=orch.commit_staged_batch(stage['batch_uuid'])
    assert result['new_records']==120 and len(queries)==1
    with get_connection() as c:assert [r[0] for r in c.execute('SELECT DISTINCT symbol FROM indicator_cache')]==['UNTOUCHED']

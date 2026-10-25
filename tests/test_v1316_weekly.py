"""Synthetic-only account-week boundaries, exports and presentation regressions."""
from decimal import Decimal
from pathlib import Path
import pytest
from psx_app.database.schema import init_db, get_connection
from psx_app.services.account_service import AccountService

ROOT = Path(__file__).resolve().parents[1]

@pytest.fixture
def weekly_db(tmp_path):
    path = tmp_path/'weekly.db'
    init_db(path)
    with get_connection(path) as c:
        for account, dates in [('91404',['2020-12-31','2021-01-01','2021-01-04','2021-01-05']),('91269',['2021-01-01'])]:
            for i, dt in enumerate(dates,1):
                c.execute('''INSERT INTO account_daily_symbol_summary
                    (account_id,trade_date,symbol,buy_qty,buy_value,sell_qty,sell_value,
                     commission_amount,sst_amount,updated_at)
                    VALUES (?,?,?,?,?,?,?,?,?,?)''',
                    (account,dt,'TEST',i*10,i*1100,i*10,i*1200,i*.1,i*.013,'2026-01-01'))
    return path

def test_iso_week_crosses_calendar_year_without_splitting(weekly_db):
    rep = AccountService(weekly_db).get_account_report('91404')
    weeks = rep['weekly_summary']['rows']
    assert [w['week_label'] for w in weeks] == ['2020-W53','2021-W01']
    assert weeks[0]['week_start']=='2020-12-28' and weeks[0]['week_end']=='2021-01-03'
    assert weeks[0]['first_date']=='2020-12-31' and weeks[0]['last_date']=='2021-01-01'
    assert weeks[0]['trading_days']==2 and weeks[0]['buy_qty']==30
    assert weeks[1]['week_start']=='2021-01-04' and weeks[1]['trading_days']==2
    assert weeks[0]['avg_buy_price']==110 and weeks[0]['avg_sell_price']==120
    assert weeks[0]['net_pnl']==pytest.approx(300-.339)
    assert weeks[0]['return_on_buy']==pytest.approx((300-.339)/3300)

def test_weekly_totals_reconcile_and_accounts_remain_separate(weekly_db):
    svc=AccountService(weekly_db)
    a=svc.get_account_report('91404');b=svc.get_account_report('91269')
    for rep in [a,b]:
        weekly=rep['weekly_summary']
        for field in ['buy_qty','sell_qty','buy_value','sell_value','commission_amount','sst_amount','total_charges','gross_pnl','net_pnl']:
            assert sum(Decimal(str(w[field])) for w in weekly['rows']) == pytest.approx(Decimal(str(weekly['total'][field])))
            assert weekly['total'][field]==rep['monthly_summary']['total'][field]==rep['yearly_summary']['total'][field]
    assert b['weekly_summary']['total']['buy_qty']==10
    assert a['weekly_summary']['total']['buy_qty']==100
    assert len(b['weekly_summary']['rows'])==1
    with pytest.raises(ValueError):svc.get_account_report('99999')

def test_empty_weekly_report_has_no_fabricated_totals(tmp_path):
    p=tmp_path/'empty.db';init_db(p)
    rep=AccountService(p).get_account_report('91404')['weekly_summary']
    assert rep['rows']==[] and rep['total'] is None

def test_excel_weekly_sheet_uses_same_values_and_freezes_headers(weekly_db,tmp_path):
    import openpyxl
    from psx_app.services.export_service import ExportService
    path=ExportService(db_path=weekly_db,data_root=tmp_path/'exports-root').export_account_workbook('91404')
    book=openpyxl.load_workbook(path,data_only=True)
    ws=book['Weekly Summary']
    assert ws.freeze_panes=='B5' and ws['B5'].value==30
    assert '2020-W53' in ws['A5'].value
    assert 'THM00601' not in book['Dashboard']['A2'].value
    assert ws['B7'].value==100  # TOTAL from all four days, not averages of weeks.

def test_ui_weekly_tables_typography_and_consolidation():
    app=(ROOT/'psx_app/static/app.js').read_text();ui=(ROOT/'psx_app/static/upgrade139.js').read_text();css=(ROOT/'psx_app/static/styles.css').read_text()
    assert '["weekly_summary", "Weekly Summary"]' in app
    assert 'Performance (THM00601)' not in app
    nav=ui[ui.index('function techNav139'):ui.index('function skeleton139')]
    assert 'Stock screener' in nav and 'Which stocks have data?' not in nav
    assert "tab==='all531'&&id==='core36'" in nav
    assert 'id="screenerScope"' in app and 'Data diagnostics' in ui
    assert 'loadLearningOnDemand' in ui and 'ontoggle="loadLearningOnDemand(this)"' in ui
    initial=ui[ui.index('const [an,contract]'):ui.index('if(gen!==analystGeneration')]
    assert '/learning/' not in initial
    assert 'max-height:clamp(260px,60vh,620px)' in css and 'thead{position:sticky;top:0' in css
    assert 'data-view="futures"] .data-table td{font-size:13.5px}' in css
    assert '@media print{.table-wrap{max-height:none!important' in css

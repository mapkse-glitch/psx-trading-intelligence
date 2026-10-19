import openpyxl
import pytest
from psx_app.parsers.ohlcv_parser import split_futures_contract_symbol
from psx_app.services.export_service import ExportService
from psx_app.services.futures_service import FuturesService


def test_dynamic_contract_symbol_splitting():
    assert split_futures_contract_symbol("ATRL-OCTB") == ("ATRL", "OCTB")
    assert split_futures_contract_symbol("OGDC-SEP") == ("OGDC", "SEP")
    assert split_futures_contract_symbol("BOP-SEPB") == ("BOP", "SEPB")
    assert split_futures_contract_symbol("NML-AUG") == ("NML", "AUG")


def test_futures_comparison_formulas_sorting_and_days_recalculation():
    svc = FuturesService()
    comp30 = svc.get_comparison_rows("monthly_comparison", selected_date="2026-10-16", days=30)
    assert comp30["rows_count"] == 98

    # First row must be lowest Annualized % (KAPCO-OCTB on 2026-10-16)
    first = comp30["rows"][0]
    assert first["symbol"] == "KAPCO"
    assert first["contract"] == "OCTB"
    assert pytest.approx(first["current_value"], abs=1e-4) == 26.54
    assert pytest.approx(first["future_value"], abs=1e-4) == 25.50
    assert pytest.approx(first["difference"], abs=1e-4) == -1.04
    expected_diff_pct = (-1.04 / 26.54) * 100.0
    expected_ann_pct = (-1.04 / 26.54) * (365.0 / 30.0) * 100.0
    assert pytest.approx(first["difference_pct"], abs=1e-4) == expected_diff_pct
    assert pytest.approx(first["annualized_pct"], abs=1e-4) == expected_ann_pct

    # Verify ascending sort across all valid rows
    ann_vals = [r["annualized_pct"] for r in comp30["rows"] if isinstance(r["annualized_pct"], (int, float))]
    assert ann_vals == sorted(ann_vals)

    # Verify dynamic recalculation when Days changes from 30 to 15 (Annualized % doubles)
    comp15 = svc.get_comparison_rows("monthly_comparison", selected_date="2026-10-16", days=15)
    assert pytest.approx(comp15["rows"][0]["annualized_pct"], abs=1e-4) == expected_ann_pct * 2.0


def test_futures_summaries_and_rolling_window():
    svc = FuturesService()
    sym_sum = svc.get_future_symbol_summary("AGHA-OCT", apply_rolling=True)
    assert sym_sum["rows_count"] == 23
    for r in sym_sum["rows"]:
        assert pytest.approx(r["Change"], abs=1e-6) == r["CLOSE"] - r["PR_CL"]

    # Symbol Filter must return all 20 retained dates for OCTB + OGDC (2026-09-21 to 2026-10-16)
    sym_filt = svc.get_comparison_rows("symbol_filter", contract_month="OCTB", base_symbol="OGDC", apply_rolling=True)
    assert sym_filt["rows_count"] == 20
    assert sym_filt["rows"][0]["trade_date"] == "2026-09-21"
    assert sym_filt["rows"][-1]["trade_date"] == "2026-10-16"

    # Month Filter must return 56 rows for 2026-10-16 + OCT
    mon_filt = svc.get_comparison_rows("month_filter", selected_date="2026-10-16", contract_month="OCT", apply_rolling=True)
    assert mon_filt["rows_count"] == 56

    date_sum = svc.get_date_summary("2026-10-16", apply_rolling=True)
    assert date_sum["rows_count"] == 98
    syms = [r["SYMBOL"] for r in date_sum["rows"]]
    assert syms == sorted(syms)

    # Non-existent date returns empty rows with message
    empty_sum = svc.get_date_summary("2026-12-31", apply_rolling=True)
    assert empty_sum["rows_count"] == 0
    assert empty_sum["message"] == "No data available for this date."

    prev = svc.preview_rolling_window_cleanup(months=2)
    assert prev["cutoff_date"] == "2026-09-01"
    assert prev["futures_rows_retained_in_active_window"] == 3913


def test_futures_8_sheet_excel_export_structure():
    exp = ExportService()
    p = exp.export_futures_workbook()
    wb = openpyxl.load_workbook(p, read_only=False)
    expected_sheets = [
        "Monthly Comparison",
        "Month Filter",
        "Symbol Filter",
        "Annualized Comparison",
        "Date Filter",
        "Future Symbol Summary",
        "Date Summary",
        "Data",
    ]
    assert wb.sheetnames == expected_sheets
    visible = [s for s in wb.sheetnames if wb[s].sheet_state == "visible"]
    hidden = [s for s in wb.sheetnames if wb[s].sheet_state == "hidden"]
    assert len(visible) == 7
    assert hidden == ["Data"]
    wb.close()

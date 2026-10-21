"""Synthetic, portable regressions for the 1.3 audit corrections."""
import base64
import gzip
import json
import sqlite3
from datetime import datetime, timedelta
from pathlib import Path
import openpyxl
import pytest
from fastapi.testclient import TestClient
from psx_app.main import app, _parse_sheet_filters
from psx_app.database.schema import init_db, get_connection
from psx_app.services.backup_service import BackupService
from psx_app.services.cloud_sync_service import CloudSyncService, _LAST_SYNC_RESULT
from psx_app.services.data_guard import date_integrity
from psx_app.services.export_service import ExportService
from psx_app.services.intelligence_service import IntelligenceService
from psx_app.services.reconciliation_service import ReconciliationService
from psx_app.services.technical_service import TechnicalService

@pytest.fixture
def fresh_db(tmp_path):
    db = tmp_path / 'database' / 'app.db'
    db.parent.mkdir()
    init_db(db)
    return db


def add_market(db):
    with get_connection(db) as c:
        for date in ('2026-10-12','2026-10-13','2026-10-16'):
            c.execute('INSERT INTO equity_ohlcv(trade_date,symbol,close,open,high,low,pr_close,volume) VALUES (?,?,?,?,?,?,?,?)', (date,'OGDC',100,99,102,98,99,1000))
            for month,close in [('OCT',101),('OCTB',102)]:
                c.execute('INSERT INTO futures_ohlcv(trade_date,contract_symbol,base_symbol,contract_month,close,open,high,low,pr_close,volume,avg_price,change_val) VALUES (?,?,?,?,?,?,?,?,?,?,?,?)', (date,'OGDC-'+month,'OGDC',month,close,100,104,99,100,500,101,close-100))
        c.commit()


def owner_client(monkeypatch):
    monkeypatch.delenv('PSX_AUTH_DISABLED', raising=False)
    monkeypatch.setenv('PSX_ADMIN_PASSWORD', 'test-only-long-password')
    monkeypatch.setenv('PSX_SESSION_SECRET', 'test-only-signing-key-' + 'x'*40)
    return TestClient(app)


def test_auth_fail_closed_and_private_endpoints(monkeypatch):
    monkeypatch.delenv('PSX_AUTH_DISABLED', raising=False)
    monkeypatch.delenv('PSX_ADMIN_PASSWORD', raising=False)
    monkeypatch.delenv('PSX_SESSION_SECRET', raising=False)
    client = TestClient(app)
    assert client.get('/').status_code == 200
    assert client.get('/api/auth/status').json()['configured'] is False
    for path in ['/api/accounts/91404','/api/settings','/api/health','/api/cloud-sync/download-snapshot','/api/exports/generate/futures_xlsx','/openapi.json','/docs']:
        assert client.get(path).status_code == 401, path
    assert client.post('/api/auth/login', json={'password':'anything'}).status_code == 503


def test_login_csrf_logout_and_cookie_flags(monkeypatch):
    client = owner_client(monkeypatch)
    assert client.post('/api/auth/login', json={'password':'incorrect'}).status_code == 401
    login = client.post('/api/auth/login', json={'password':'test-only-long-password','remember':True})
    assert login.status_code == 200
    assert 'HttpOnly' in login.headers['set-cookie']
    assert 'SameSite=strict' in login.headers['set-cookie']
    assert 'Max-Age=604800' in login.headers['set-cookie']
    status = client.get('/api/auth/status').json()
    assert status['authenticated'] and status['csrf']
    assert client.get('/api/settings').status_code == 200
    assert client.post('/api/backups/create',json={}).status_code == 403
    assert client.post('/api/settings',json={'updates':{}},headers={'X-PSX-CSRF':status['csrf'], 'Origin':'https://evil.example'}).status_code == 403
    assert client.post('/api/auth/logout',headers={'X-PSX-CSRF':status['csrf']}).status_code == 200
    assert client.get('/api/settings').status_code == 401


def test_serverless_cannot_enable_local_auth_bypass(monkeypatch):
    monkeypatch.setenv('VERCEL','1')
    monkeypatch.setenv('PSX_AUTH_DISABLED','1')
    monkeypatch.delenv('PSX_ADMIN_PASSWORD',raising=False)
    assert TestClient(app).get('/api/settings').status_code == 401


def test_auth_cookie_tamper_expiry_and_cross_origin(monkeypatch):
    client = owner_client(monkeypatch)
    assert client.post('/api/auth/login',json={'password':'test-only-long-password'}, headers={'Origin':'https://evil.example'}).status_code == 403
    from psx_app.security import COOKIE, sign
    payload=base64.urlsafe_b64encode(json.dumps({'exp':1,'csrf':'bad'}).encode()).decode()
    client.cookies.set(COOKIE,payload+'.'+sign(payload))
    assert client.get('/api/settings').status_code == 401
    client.cookies.set(COOKIE,payload+'.'+'0'*64)
    assert client.get('/api/settings').status_code == 401


class FakeGit:
    def __init__(self, db):
        self.blobs={'initial':base64.b64encode(gzip.compress(db.read_bytes())).decode()}
        self.head='head0';self.tree='tree0';self.blob='initial';self.counter=0;self.commits={};self.trees={}
    def request(self, method, endpoint, payload=None):
        if endpoint == "": return {"private": True}
        if method=='GET' and endpoint.startswith('/git/ref/'):
            return {'object':{'sha':self.head}}
        if method=='GET' and endpoint.startswith('/git/commits/'):
            return {'tree':{'sha':self.tree}}
        if method=='GET' and endpoint.startswith('/git/trees/'):
            return {'tree':[{'path':'latest_cloud_sync.db.gz','sha':self.blob}]}
        if method=='GET' and endpoint.startswith('/git/blobs/'):
            return {'content':self.blobs[endpoint.split('/')[-1]]}
        self.counter+=1; sha='sha'+str(self.counter)
        if endpoint=='/git/blobs':self.blobs[sha]=payload['content']
        elif endpoint=='/git/trees':self.trees[sha]=payload['tree'][0]['sha']
        elif endpoint=='/git/commits':self.commits[sha]=payload
        elif method=='PATCH':
            commit=self.commits[payload['sha']]
            if commit['parents'] != [self.head]:raise RuntimeError('409 ref advanced')
            self.head=payload['sha'];self.tree=commit['tree'];self.blob=self.trees[self.tree]
        return {'sha':sha}


def cloud(db, remote, monkeypatch):
    service=CloudSyncService(db);service.repo='private/repo';service.token='test-token'
    monkeypatch.setattr(service,'_gh_request',remote.request)
    return service


def test_stale_worker_cannot_overwrite_newer_snapshot(fresh_db,tmp_path,monkeypatch):
    remote=FakeGit(fresh_db)
    a=cloud(tmp_path/'a'/'database'/'app.db', remote, monkeypatch)
    b=cloud(tmp_path/'b'/'database'/'app.db', remote, monkeypatch)
    assert a.pull_latest_db_on_startup()['pulled']
    assert b.pull_latest_db_on_startup()['pulled']
    a.assert_current_base()
    with sqlite3.connect(a.db_path) as c:c.execute('CREATE TABLE owner_marker (value TEXT)');c.execute("INSERT INTO owner_marker VALUES ('A')")
    a.mark_pending()
    assert a.push_db_after_update('A')['pushed']
    with pytest.raises(RuntimeError,match='stale'):
        b.assert_current_base()
    with sqlite3.connect(b.db_path) as c:c.execute('CREATE TABLE owner_marker (value TEXT)');c.execute("INSERT INTO owner_marker VALUES ('B')")
    assert not b.push_db_after_update('B')['pushed']
    assert b.base_state()['pending']
    assert not b.pull_latest_db_on_startup()['pulled']
    saved=tmp_path/'remote.db';saved.write_bytes(gzip.decompress(base64.b64decode(remote.blobs[remote.blob])))
    with sqlite3.connect(saved) as c:assert c.execute('SELECT value FROM owner_marker').fetchall()==[('A',)]


def test_unknown_base_and_unconfigured_serverless_refuse_writes(fresh_db,monkeypatch):
    svc=CloudSyncService(fresh_db);svc.token='test-token';svc.repo='private/repo'
    monkeypatch.setattr(svc,'_remote_head',lambda:('head','tree','existing'))
    with pytest.raises(RuntimeError,match='baseline'):svc.assert_current_base()
    svc.token='';monkeypatch.setenv('VERCEL','1')
    with pytest.raises(RuntimeError,match='blocked'):svc.assert_current_base()


def test_sync_failure_does_not_claim_cloud_success(fresh_db,monkeypatch):
    _LAST_SYNC_RESULT.clear()
    svc=CloudSyncService(fresh_db);svc.token='test-token';svc.repo='private/repo'
    monkeypatch.setattr(svc,'_gh_request',lambda *args,**kwargs: (_ for _ in ()).throw(RuntimeError('offline')))
    result=svc.push_db_after_update('offline')
    assert not result['pushed'] and result['local_snapshot_updated']
    status=svc.get_status()
    assert status['health']=='PENDING' and status['last_push'] is None
    assert 'offline' in status['last_error']


def test_backups_include_wal_and_never_silently_delete(fresh_db,tmp_path):
    svc=BackupService(fresh_db,tmp_path)
    conn=sqlite3.connect(fresh_db)
    conn.execute('PRAGMA journal_mode=WAL')
    conn.execute('CREATE TABLE wal_marker (value TEXT)');conn.execute("INSERT INTO wal_marker VALUES ('committed')");conn.commit()
    import zipfile
    first=svc.create_backup('wal')
    with zipfile.ZipFile(first['zip_path']) as zf:
        restored=tmp_path/'verify.db';restored.write_bytes(zf.read('app.db'))
    with sqlite3.connect(restored) as check:assert check.execute('SELECT value FROM wal_marker').fetchone()[0]=='committed'
    for _ in range(21):svc.create_backup('retained')
    assert len(svc.list_backups())==22
    conn.close()
    assert svc.restore_backup(first['zip_filename'])['post_restore_integrity']['ok']


def test_restore_endpoint_requests_cloud_save(monkeypatch):
    calls=[]
    monkeypatch.setattr(BackupService,'restore_backup',lambda self,name:{'status':'RESTORED'})
    monkeypatch.setattr(CloudSyncService,'push_db_after_update',lambda self,reason: calls.append(reason) or {'pushed':True})
    response=TestClient(app).post('/api/backups/restore',json={'zip_filename':'test.zip'})
    assert response.status_code==200 and response.json()['cloud_sync']['pushed']
    assert calls==['Backup restored']


def test_futures_export_respects_independent_filters(fresh_db,tmp_path):
    add_market(fresh_db)
    filters={'monthly_comparison':{'selected_date':'2026-10-12','days':10},
             'month_filter':{'selected_date':'2026-10-13','days':20,'contract_month':'OCT'},
             'symbol_filter':{'selected_date':'2026-10-16','days':40,'contract_month':'OCT','base_symbol':'OGDC'},
             'annualized_comparison':{'selected_date':'2026-10-18','days':50},
             'date_filter':{'selected_date':'2026-10-13','days':60},
             'future_symbol_summary':{'contract_symbol':'OGDC-OCTB'},
             'date_summary':{'selected_date':'2026-10-12'},'apply_rolling':False}
    path=ExportService(fresh_db,tmp_path).export_futures_workbook(sheet_filters=filters)
    wb=openpyxl.load_workbook(path)
    assert len(wb.sheetnames)==8 and wb['Data'].sheet_state=='hidden'
    assert wb['Monthly Comparison']['C4'].value.date().isoformat()=='2026-10-12'
    assert wb['Month Filter']['C4'].value.date().isoformat()=='2026-10-13'
    assert wb['Monthly Comparison']['F4'].value==10
    assert wb['Month Filter']['F4'].value==20
    sf=wb['Symbol Filter']
    assert sf['C4'].value=='OCT' and sf['F4'].value==40
    assert sf.max_row==9 and {sf.cell(r,3).value for r in range(7,10)}=={'OCT'}
    assert wb['Annualized Comparison']['C4'].value.date().isoformat()=='2026-10-16'
    assert '18-10-2026' in wb['Annualized Comparison']['A3'].value
    assert wb['Date Summary']['C4'].value.date().isoformat()=='2026-10-12'
    assert wb['Future Symbol Summary']['C4'].value=='OGDC-OCTB'
    assert sf['A7'].number_format=='DD-MM-YYYY'
    wb.close()


@pytest.mark.parametrize('value',['[]','{"unknown":{}}','{"month_filter":{"days":0}}','{"month_filter":{"days":true}}','{"month_filter":{"selected_date":"bad"}}','{"apply_rolling":"true"}'])
def test_workbook_filter_validation(value):
    from fastapi import HTTPException
    with pytest.raises(HTTPException):_parse_sheet_filters(value)


def test_future_dated_history_blocks_setup_generation(fresh_db):
    date=(datetime.now()+timedelta(days=365)).date().isoformat()
    with get_connection(fresh_db) as conn:
        conn.execute('INSERT INTO equity_ohlcv(trade_date,symbol,close) VALUES (?,?,?)',(date,'ATRL',100));conn.commit()
    assert not date_integrity(fresh_db)['ok']
    svc=IntelligenceService(fresh_db)
    assert svc.generate_trading_ideas()['status']=='BLOCKED'
    assert svc.run_automated_technical_analyst()['status']=='BLOCKED'


def test_missing_references_are_not_checked(fresh_db,tmp_path):
    result=ReconciliationService(fresh_db,tmp_path/'no-fixtures').reconcile_against_legacy_workbooks()
    assert len(result)==5 and all(r['status']=='NOT_CHECKED' for r in result.values())


def test_atrl_parity_detects_changed_close(fresh_db,tmp_path):
    start=datetime(2025,1,1)
    with get_connection(fresh_db) as c:
        for i in range(220):
            close=100+i*.2
            c.execute('INSERT INTO equity_ohlcv(trade_date,symbol,open,high,low,close,pr_close,volume) VALUES (?,?,?,?,?,?,?,?)', ((start+timedelta(days=i)).date().isoformat(),'ATRL',close,close+2,close-2,close,close-.2,1000))
        c.commit()
    new=TechnicalService(fresh_db).get_ticker_analysis('ATRL')['latest']
    out=tmp_path/'refs'/'Output folder';out.mkdir(parents=True)
    wb=openpyxl.Workbook();ws=wb.active;ws.title='ATRL_DATA';ws.append(['date'])
    row=[None]*21;row[0]=datetime.fromisoformat(new['latest_date'])
    for idx,key in {4:'close',6:'rsi14',7:'macd',8:'macd_sig',9:'macd_hist',10:'bb_up',11:'bb_mid',12:'bb_low',15:'ema9',16:'ema21',17:'ema50',18:'ema200_recursive',20:'adx14'}.items():row[idx]=new[key]
    ws.append(row);wb.save(out/'PSX_RSI_Pivot_Analysis (4).xlsx');wb.close()
    service=ReconciliationService(fresh_db,tmp_path/'refs')
    assert service.reconcile_against_legacy_workbooks()['module_d_technical_atrl']['status']=='MATCH'
    with get_connection(fresh_db) as c:c.execute('UPDATE equity_ohlcv SET close=close+10 WHERE trade_date=?',(new['latest_date'],));c.commit()
    result=service.reconcile_against_legacy_workbooks()['module_d_technical_atrl']
    assert result['status']=='DISCREPANCY' and 'close' in result['mismatches']


@pytest.mark.parametrize('updates',[{'sst_rate':-1},{'sst_rate':2},{'futures_default_days':0},{'futures_default_days':1.5},{'unexpected':3},{'commission_slab1_upper':100}])
def test_settings_reject_invalid_values(updates):
    response=TestClient(app).post('/api/settings',json={'updates':updates})
    assert response.status_code==422


def test_synthetic_import_is_idempotent(fresh_db,tmp_path):
    from psx_app.services.import_orchestrator import ImportOrchestrator
    service=ImportOrchestrator(fresh_db,tmp_path)
    content=b'SYN Price Summary for 16/10/2026 as per 16/10/2026, 15:47:20\n"100.00","2","1000","100000.00"\n'
    first=service.stage_uploaded_files([('SYN.txt',content)])
    result=service.commit_staged_batch(first['batch_uuid'])
    assert result['new_records']==1
    again=service.stage_uploaded_files([('SYN.txt',content)])
    second=service.commit_staged_batch(again['batch_uuid'])
    assert second['new_records']==0
    with get_connection(fresh_db) as c:assert c.execute('SELECT COUNT(*) FROM top5_volume_records').fetchone()[0]==1


def test_public_repository_sync_is_refused(fresh_db,monkeypatch):
    svc=CloudSyncService(fresh_db);svc.token='test-token';svc.repo='public/repo'
    monkeypatch.setattr(svc,'_gh_request',lambda *args,**kwargs:{'private':False})
    with pytest.raises(RuntimeError,match='private GitHub'):
        svc.assert_current_base()


def test_date_repair_requires_review_and_rejects_stale_plan(fresh_db,tmp_path):
    from psx_app.services.date_repair_service import DateRepairService
    service=DateRepairService(fresh_db,tmp_path)
    plan=service.plan()
    with get_connection(fresh_db) as c:
        c.execute("INSERT INTO equity_ohlcv(trade_date,symbol,close) VALUES ('2026-10-16','ATRL',100)");c.commit()
    with pytest.raises(ValueError,match='changed since'):
        service.apply(expected_plan_token=plan['plan_token'])
    assert len(BackupService(fresh_db,tmp_path).list_backups())==1
    response=TestClient(app).post('/api/admin/date-normalization/apply',json={'confirm':True})
    assert response.status_code==400 and 'plan_token' in response.json()['detail']


def test_repair_backup_is_verified_restorable_zip(fresh_db,tmp_path):
    from psx_app.services.date_repair_service import DateRepairService
    service=DateRepairService(fresh_db,tmp_path)
    result=service.apply(expected_plan_token=service.plan()['plan_token'])
    path=Path(result['backup_path'])
    assert path.suffix=='.zip' and path.exists()
    backups=BackupService(fresh_db,tmp_path)
    assert path.name in {b['zip_filename'] for b in backups.list_backups()}
    assert backups.restore_backup(path.name)['post_restore_integrity']['ok']


def test_failed_multi_file_chunk_stage_preserves_completed_parts():
    client=TestClient(app)
    content=b'SYN Price Summary for 16/10/2026 as per 16/10/2026, 15:47:20\n"100.00","2","1000","100000.00"\n'
    one='release-test-file-one';two='release-test-file-two'
    assert client.post('/api/imports/upload-chunk',params={'upload_id':one,'name':'SYN.txt','idx':0,'total':1},content=content).status_code==200
    request={'uploads':[{'upload_id':one,'name':'SYN.txt','total_chunks':1},{'upload_id':two,'name':'SYN2.txt','total_chunks':1}]}
    assert client.post('/api/imports/stage-chunked',json=request).status_code==409
    from psx_app.main import _chunk_dir
    assert (_chunk_dir(one)/'000000.part').exists()
    assert client.post('/api/imports/upload-chunk',params={'upload_id':two,'name':'SYN2.txt','idx':0,'total':1},content=content).status_code==200
    assert client.post('/api/imports/stage-chunked',json=request).status_code==200
    assert not _chunk_dir(one).exists()


def test_empty_workspace_and_exports_are_graceful(fresh_db,tmp_path,monkeypatch):
    import importlib
    main=importlib.import_module('psx_app.main')
    monkeypatch.setattr(main,'DB_PATH',fresh_db);monkeypatch.setattr(main,'DATA_ROOT',tmp_path)
    main.invalidate_api_cache()
    client=TestClient(app,raise_server_exceptions=False)
    for path in ['/api/overview','/api/system/status','/api/reconciliation','/api/technical/universe']:
        response=client.get(path)
        assert response.status_code==200,(path,response.text[:300])
    for target in ['top5_xlsx','futures_xlsx','account_91404_xlsx','account_91269_xlsx','technical_xlsx']:
        response=client.get('/api/exports/generate/'+target)
        assert response.status_code==200,(target,response.text[:300])
    main.invalidate_api_cache()


def test_setup_context_is_same_date_and_most_liquid(fresh_db,monkeypatch):
    add_market(fresh_db)
    with get_connection(fresh_db) as c:
        c.execute("UPDATE futures_ohlcv SET volume=9000 WHERE contract_month='OCTB'")
        c.execute("INSERT INTO top5_volume_records(trade_date,ticker,rank,rate,volume,trades,imported_at) VALUES ('2026-10-12','OGDC',1,100,123,1,'2026-10-12T00:00:00Z')")
        c.commit()
    item={'symbol':'OGDC','latest_date':'2026-10-16','bars_count':50,'close':100,'rsi14':55,'rsi_zone':'NEUTRAL',
          'signal':{'score':4,'combined_signal':'BUY'},'ema9':101,'ema21':99,'ema50':98,
          'daily_pivots':{'classic':{'P':100,'R1':102,'R2':104,'R3':106,'S1':98,'S2':96}},'is_futures_eligible':True}
    svc=IntelligenceService(fresh_db)
    monkeypatch.setattr(svc.tech_service,'get_universe_summary',lambda **kw:{'items':[item]})
    result=svc.generate_trading_ideas()
    setup=result['setups'][0]
    assert 'probability_score' not in setup and 'not a probability' in setup['score_interpretation']
    notes=' '.join(setup['fundamental_why'])
    assert 'OGDC-OCTB close 102.00' in notes and 'Top 5 register context' not in notes
    assert svc.generate_trading_ideas(scope_or_sector='TOP5_VOL')['setups_count']==0
    item['latest_date']='2026-10-12'
    matched=svc.generate_trading_ideas(scope_or_sector='TOP5_VOL')['setups'][0]
    assert '123 shares' in ' '.join(matched['fundamental_why'])


def test_https_cookie_and_malformed_login(monkeypatch):
    owner_client(monkeypatch)
    client=TestClient(app,base_url='https://testserver')
    assert client.post('/api/auth/login',json=[]).status_code==400
    response=client.post('/api/auth/login',json={'password':'test-only-long-password','remember':True})
    assert response.status_code==200
    assert 'Secure' in response.headers['set-cookie']
    assert client.get('/api/settings').status_code==200

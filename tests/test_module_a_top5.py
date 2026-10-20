import openpyxl
from pathlib import Path
from psx_app.parsers.top5_parser import parse_top5_txt
from psx_app.services.top5_service import Top5Service
from psx_app.services.export_service import ExportService


def test_top5_single_txt_parsing_and_volume_ranking():
    txt_path = Path("/home/user/drive_mirror/Input folder/ATRL.txt")
    content = txt_path.read_text(encoding="utf-8-sig")
    res = parse_top5_txt(content, "ATRL.txt")

    assert res["ticker"] == "ATRL"
    assert res["trade_date"] == "2026-10-16"
    assert res["mmddyyyy"] == "10/16/2026"
    assert len(res["top5"]) == 5

    # Exact verification against legacy workbook 10/16/2026 block
    r1 = res["top5"][0]
    assert r1["rank"] == 1
    assert r1["rate"] == 1215.0
    assert r1["volume"] == 89051
    assert r1["trades"] == 220

    # Ensure strictly non-increasing volume order
    vols = [r["volume"] for r in res["top5"]]
    assert vols == sorted(vols, reverse=True)


def test_top5_fewer_than_5_rows_and_invalid_row_rejection():
    sample = (
        "\ufeffTEST Price Summary for 16/10/2026 as per 16/10/2026, 15:47:20\n"
        '"100.50","5","2,000","201,000.00"\n'
        '"101.00","12","5,500","555,500.00"\n'
        '"-10.00","2","1,000","-10,000.00"\n'
        '"102.00","0","0","0.00"\n'
        '"corrupted_line_here"\n'
        '"99.25","3","800","79,400.00"\n'
    )
    res = parse_top5_txt(sample, "TEST.txt")
    assert res["ticker"] == "TEST"
    assert res["trade_date"] == "2026-10-16"
    # Only 3 valid rows exist; parser must NOT invent rows 4 or 5
    assert len(res["top5"]) == 3
    assert [r["rank"] for r in res["top5"]] == [1, 2, 3]
    assert [r["volume"] for r in res["top5"]] == [5500, 2000, 800]
    # Issues logged for negative rate, zero volume/trades, and malformed row + fewer than 5 rows
    issue_types = {i["issue_type"] for i in res["issues"]}
    assert "NON_POSITIVE_VALUES" in issue_types
    assert "FEWER_THAN_5_ROWS" in issue_types


def test_top5_tie_volume_detection():
    sample = (
        "TIE Price Summary for 16/10/2026 as per 16/10/2026, 15:47:20\n"
        '"50.00","10","5,000","250,000.00"\n'
        '"51.00","15","5,000","255,000.00"\n'
    )
    res = parse_top5_txt(sample, "TIE.txt")
    assert len(res["top5"]) == 2
    # Higher trades breaks tie deterministically
    assert res["top5"][0]["trades"] == 15
    assert any(i["issue_type"] in ("VOLUME_TIE", "TIED_VOLUMES") for i in res["issues"])


def test_top5_service_and_legacy_excel_export():
    svc = Top5Service()
    reg = svc.get_register(ticker="ATRL", start_date="2026-10-16", end_date="2026-10-16")
    assert reg["blocks_count"] == 1
    blk = reg["date_blocks"][0]
    assert blk["title"] == "ATRL — Top 5 Volume Trades (16-10-2026)"
    assert len(blk["rows"]) == 5

    exp = ExportService()
    xlsx_path = exp.export_top5_workbook()
    assert xlsx_path.exists()
    wb = openpyxl.load_workbook(xlsx_path, read_only=False)
    assert "ATRL" in wb.sheetnames
    ws = wb["ATRL"]
    # Verify headers are strictly Rank, Rate, Volume, Trades
    assert [ws.cell(3, c).value for c in range(1, 5)] == ["Rank", "Rate", "Volume", "Trades"]
    assert ws.cell(3, 5).value is None
    assert len(ws.tables) >= 40
    wb.close()

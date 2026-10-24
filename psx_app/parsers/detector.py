import re
from pathlib import Path
from typing import Dict, Any, Optional
import openpyxl


def _detect_price_sheet(ws) -> Optional[str]:
    """Recognise the active raw-price sheet by columns, never by its year."""
    rows = list(ws.iter_rows(max_row=32, values_only=True))
    for header_index, row in enumerate(rows[:5]):
        names = [str(v or "").strip().upper() for v in row]
        columns = {name: i for i, name in enumerate(names) if name}
        if not {"DATE", "SYMBOL", "OPEN", "HIGH", "LOW", "CLOSE"}.issubset(columns):
            continue
        if not {"VOLUME", "VOL"}.intersection(columns):
            continue
        if not {"PR-CLOSE", "PR_CL", "PR_CLOSE", "PREV_CLOSE"}.intersection(columns):
            continue
        symbols = [str(r[columns["SYMBOL"]]).strip().upper()
                   for r in rows[header_index + 1:]
                   if len(r) > columns["SYMBOL"] and r[columns["SYMBOL"]] is not None]
        if not symbols:
            return None
        is_contract = [bool(re.search(
            r"[-\s](?:JAN|FEB|MAR|APR|MAY|JUN|JUL|AUG|SEP|OCT|NOV|DEC)[A-Z0-9]*$", s
        )) for s in symbols]
        if all(is_contract):
            return "OHLCV_FUTURES"
        if not any(is_contract):
            return "OHLCV_CASH"
        return None  # Do not silently put a mixed cash/futures sheet in one table.
    return None


def detect_file_type(file_path: Path, original_filename: Optional[str] = None) -> Dict[str, Any]:
    """Inspect file header and structure to determine canonical PSX source type."""
    fname = original_filename or file_path.name
    lower = fname.lower()

    if lower.endswith(".txt") or lower.endswith(".csv"):
        try:
            text_head = file_path.read_text(encoding="utf-8-sig", errors="replace")[:4096]
        except Exception as e:
            return {"file_type": "UNKNOWN", "filename": fname, "error": str(e)}

        lines = [ln.strip() for ln in text_head.splitlines() if ln.strip()]
        first = lines[0] if lines else ""

        if "Price Summary for" in first:
            m = re.match(
                r"^([A-Za-z0-9_-]+)\s+Price Summary for\s+(\d{2}/\d{2}/\d{4})",
                first,
            )
            ticker = m.group(1).upper() if m else Path(fname).stem.upper()
            return {
                "file_type": "TOP5_TXT",
                "module": "MODULE_A",
                "filename": fname,
                "ticker": ticker,
                "header": first,
            }

        if "Daily Trade Log" in first or lower.startswith("tradelog"):
            return {
                "file_type": "TRADELOG",
                "module": "MODULE_B",
                "filename": fname,
                "header": first,
            }

        # Check CSV headers for 2026F or 2026_compare or TradeLog
        upper_first = first.upper()
        if "PR-CLOSE" in upper_first or "PR_CL" in upper_first or ("SYMBOL" in upper_first and "CLOSE" in upper_first):
            if "2026F" in fname.upper() or any("-" in ln for ln in lines[1:6]):
                return {"file_type": "OHLCV_FUTURES", "module": "MODULE_C", "filename": fname}
            return {"file_type": "OHLCV_CASH", "module": "MODULE_C_D", "filename": fname}

        # Fallback: check if body looks like 18-col TradeLog or 4-col Top5 TXT
        if len(lines) >= 1 and lines[0].count('","') >= 10:
            return {"file_type": "TRADELOG", "module": "MODULE_B", "filename": fname}
        if len(lines) >= 2 and lines[1].count('","') == 3:
            return {"file_type": "TOP5_TXT", "module": "MODULE_A", "filename": fname}

        return {"file_type": "CUSTOM_CSV_TXT", "module": "CUSTOM", "filename": fname}

    if lower.endswith(".xlsx") or lower.endswith(".xls"):
        try:
            wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
            sheets = wb.sheetnames
            price_type = _detect_price_sheet(wb.active)
            reference_sheets = []
            for sheet in wb:
                header = next(sheet.iter_rows(max_row=1, values_only=True), ())
                names = {str(v or '').strip().upper() for v in header}
                if {'SYMBOL', 'COMPANY NAME', 'SECTOR'} <= names:
                    reference_sheets.append(sheet.title)
            wb.close()
        except Exception as e:
            return {"file_type": "UNKNOWN", "filename": fname, "error": str(e)}

        if "Monthly Comparison" in sheets and "Future Symbol Summary" in sheets:
            return {"file_type": "LEGACY_FUTURES_XLSX", "module": "MODULE_C", "filename": fname, "sheets": sheets}
        if "Daily Trade Summary" in sheets and "Dashboard" in sheets:
            return {"file_type": "LEGACY_ACCOUNT_XLSX", "module": "MODULE_B", "filename": fname, "sheets": sheets}
        if "SUMMARY" in sheets and any(s.endswith("_DATA") for s in sheets):
            return {"file_type": "LEGACY_TECHNICAL_XLSX", "module": "MODULE_D", "filename": fname, "sheets": sheets}
        if "ATRL" in sheets and "PRL" in sheets and len(sheets) >= 15:
            return {"file_type": "LEGACY_TOP5_XLSX", "module": "MODULE_A", "filename": fname, "sheets": sheets}
        if sheets == ["INDEX"] or "INDEX" in fname.upper():
            return {"file_type": "INDEX_OHLCV", "module": "MODULE_D", "filename": fname, "sheets": sheets}
        if price_type:
            return {"file_type": price_type,
                    "module": "MODULE_C" if price_type == "OHLCV_FUTURES" else "MODULE_C_D",
                    "filename": fname, "sheets": sheets}

        if reference_sheets:
            return {"file_type": "SECTOR_REFERENCE", "module": "REFERENCE", "filename": fname, "sheets": sheets}
        return {"file_type": "CUSTOM_XLSX", "module": "CUSTOM", "filename": fname, "sheets": sheets}

    return {"file_type": "UNKNOWN", "module": "UNKNOWN", "filename": fname}

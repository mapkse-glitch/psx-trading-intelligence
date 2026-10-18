import csv
import io
import re
from datetime import datetime
from decimal import Decimal
from pathlib import Path
from typing import Dict, Any, List, Optional
import openpyxl
from psx_app.config import AUTHORIZED_ACCOUNTS, DEFAULT_SETTINGS


def compute_execution_charges(
    normalized_side: str,
    quantity: int,
    price: float,
    trade_value: float,
    settings: Optional[Dict[str, Any]] = None,
) -> Dict[str, Any]:
    """
    Deterministic PSX sell-side commission & SST calculation at the execution level.
    Slab 1: Sell price < 54 PKR -> Commission = Sell Quantity * 0.0175 / 100 PKR
    Slab 2: Sell price >= 54 PKR -> Commission = Sell Value * 0.000325 PKR
    SST: 13% of Commission
    """
    cfg = settings or DEFAULT_SETTINGS
    if normalized_side != "SELL":
        return {
            "commission_slab": "— (BUY)",
            "commission_amount": 0.0,
            "sst_amount": 0.0,
            "total_charges": 0.0,
        }

    slab2_start = Decimal(str(cfg.get("commission_slab2_start", 54.0)))
    slab1_paisa = Decimal(str(cfg.get("commission_slab1_paisa_per_share", 0.0175)))
    paisa_per_pkr = Decimal(str(cfg.get("paisa_per_pkr", 100.0)))
    slab2_rate = Decimal(str(cfg.get("commission_slab2_rate", 0.000325)))
    sst_rate = Decimal(str(cfg.get("sst_rate", 0.13)))

    d_price = Decimal(str(price))
    d_qty = Decimal(str(quantity))
    d_val = Decimal(str(trade_value))

    if d_price >= slab2_start:
        comm = d_val * slab2_rate
        slab_label = "0.0325% of sell value"
    else:
        comm = d_qty * slab1_paisa / paisa_per_pkr
        slab_label = "0.0175 paisa per share"

    sst = comm * sst_rate
    total = comm + sst
    return {
        "commission_slab": slab_label,
        "commission_amount": float(comm),
        "sst_amount": float(sst),
        "total_charges": float(total),
    }


def normalize_side(raw_side: str) -> Optional[str]:
    s = raw_side.strip().lower()
    if s in ("buy", "b"):
        return "BUY"
    if s in ("sell", "s", "short sell", "shortsell", "short_sell", "short"):
        return "SELL"
    return None


def parse_tradelog_date(raw: str) -> Optional[str]:
    s = str(raw).strip()
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%d-%b-%Y", "%m/%d/%Y"):
        try:
            return datetime.strptime(s[:10], fmt).strftime("%Y-%m-%d")
        except ValueError:
            continue
    return None


def parse_tradelog_content(
    content: str,
    filename: str = "TradeLog.txt",
    settings: Optional[Dict[str, Any]] = None,
) -> Dict[str, Any]:
    """
    Parse a PSX TradeLog file (18-column broker TXT format or headered CSV).
    Separates executions by account (91404, 91269) and isolates unknown accounts.
    """
    clean = content.lstrip("\ufeff")
    lines = [ln.strip() for ln in clean.splitlines() if ln.strip()]
    issues: List[Dict[str, Any]] = []
    executions: List[Dict[str, Any]] = []
    unexpected_executions: List[Dict[str, Any]] = []

    if not lines:
        return {
            "filename": filename,
            "detected_dates": [],
            "rows_read": 0,
            "executions": [],
            "unexpected_executions": [],
            "issues": [{"severity": "ERROR", "issue_type": "EMPTY_FILE", "message": f"{filename} is empty."}],
        }

    header_date: Optional[str] = None
    first = lines[0]
    m_hdr = re.search(r"Daily Trade Log for\s+(\d{2}/\d{2}/\d{4})", first, re.IGNORECASE)
    if m_hdr:
        header_date = datetime.strptime(m_hdr.group(1), "%d/%m/%Y").strftime("%Y-%m-%d")
        data_lines = lines[1:]
    else:
        m_fn = re.search(r"(\d{2}-\d{2}-\d{4}|\d{4}-\d{2}-\d{2})", filename)
        if m_fn:
            header_date = parse_tradelog_date(m_fn.group(1))
        data_lines = lines

    reader = csv.reader(io.StringIO("\n".join(data_lines)))
    all_rows = list(reader)

    # Check if row 0 is a column header row
    header_map: Optional[Dict[str, int]] = None
    start_idx = 0
    if all_rows:
        r0_lower = [c.strip().lower() for c in all_rows[0]]
        if any(k in r0_lower for k in ("account", "symbol", "quantity", "price", "transaction type", "side")):
            header_map = {name: i for i, name in enumerate(r0_lower)}
            start_idx = 1

    rows_read = 0
    detected_dates = set()
    seen_keys = set()

    for idx, row in enumerate(all_rows[start_idx:], start=start_idx + (2 if m_hdr else 1)):
        if not row or all(not c.strip() for c in row):
            continue
        rows_read += 1

        try:
            if header_map:
                def get_col(*names: str, default: str = "") -> str:
                    for n in names:
                        if n in header_map and header_map[n] < len(row):
                            return row[header_map[n]].strip()
                    return default

                terminal_id = get_col("terminal", "terminal_id", "user", default="THM00601")
                market_type = get_col("market", "instrument", "instrument type", "market_type", default="FUT")
                symbol = get_col("symbol", "scrip", "ticker").upper()
                account_id = get_col("account", "account_id", "client", "account number")
                row_date_raw = get_col("date", "trading date", "trade_date")
                trade_date = parse_tradelog_date(row_date_raw) if row_date_raw else header_date
                trade_time = get_col("time", "trade_time")
                order_id = get_col("order id", "order_id", "order")
                status = get_col("status", default="Filled")
                raw_side = get_col("transaction type", "side", "type", "buy/sell")
                qty_str = get_col("quantity", "qty", "executed quantity")
                rem_str = get_col("remaining quantity", "remaining_qty", default="0")
                ord_qty_str = get_col("order quantity", "order_qty", default=qty_str)
                price_str = get_col("trade price", "price", "rate")
                exec_id = get_col("execution id", "execution_id", "reference id", "ticket", "trade id")
                val_str = get_col("trade value", "value", "amount")
            else:
                if len(row) < 13:
                    issues.append({
                        "severity": "ERROR",
                        "issue_type": "MALFORMED_ROW",
                        "row_number": idx,
                        "message": f"Row {idx} has only {len(row)} columns (expected 18-col TradeLog format).",
                        "raw_payload": str(row),
                    })
                    continue
                terminal_id = row[0].strip()
                market_type = row[1].strip()
                symbol = row[2].strip().upper()
                account_id = row[3].strip()
                row_date_raw = row[4].strip()
                trade_date = parse_tradelog_date(row_date_raw) if row_date_raw else header_date
                trade_time = row[5].strip()
                order_id = row[6].strip()
                status = row[7].strip()
                raw_side = row[8].strip()
                qty_str = row[9].strip()
                rem_str = row[10].strip() if len(row) > 10 else "0"
                ord_qty_str = row[11].strip() if len(row) > 11 else qty_str
                price_str = row[12].strip()
                exec_id = row[14].strip() if len(row) > 14 else ""
                val_str = row[17].strip() if len(row) > 17 else ""

            if not account_id:
                issues.append({
                    "severity": "ERROR",
                    "issue_type": "MISSING_ACCOUNT",
                    "symbol": symbol,
                    "trade_date": trade_date,
                    "row_number": idx,
                    "message": f"Row {idx} is missing Account identifier.",
                    "raw_payload": str(row),
                })
                continue

            if not symbol:
                issues.append({
                    "severity": "ERROR",
                    "issue_type": "MISSING_SYMBOL",
                    "account_id": account_id,
                    "trade_date": trade_date,
                    "row_number": idx,
                    "message": f"Row {idx} is missing Symbol.",
                    "raw_payload": str(row),
                })
                continue

            if not trade_date:
                issues.append({
                    "severity": "ERROR",
                    "issue_type": "INVALID_DATE",
                    "account_id": account_id,
                    "symbol": symbol,
                    "row_number": idx,
                    "message": f"Row {idx} has missing or unparseable trading date.",
                    "raw_payload": str(row),
                })
                continue

            norm_side = normalize_side(raw_side)
            if not norm_side:
                issues.append({
                    "severity": "ERROR",
                    "issue_type": "INVALID_TRANSACTION_TYPE",
                    "account_id": account_id,
                    "symbol": symbol,
                    "trade_date": trade_date,
                    "row_number": idx,
                    "message": f"Row {idx} has unrecognized transaction type '{raw_side}'.",
                    "raw_payload": str(row),
                })
                continue

            quantity = int(float(qty_str.replace(",", "")))
            remaining_qty = int(float(rem_str.replace(",", ""))) if rem_str else 0
            order_qty = int(float(ord_qty_str.replace(",", ""))) if ord_qty_str else quantity
            price = float(price_str.replace(",", ""))
            if val_str:
                trade_value = float(val_str.replace(",", ""))
            else:
                trade_value = float(Decimal(str(quantity)) * Decimal(str(price)))
                issues.append({
                    "severity": "INFO",
                    "issue_type": "DERIVED_TRADE_VALUE",
                    "account_id": account_id,
                    "symbol": symbol,
                    "trade_date": trade_date,
                    "row_number": idx,
                    "message": f"Row {idx} had blank trade_value; derived {trade_value:.2f} from Qty*Price.",
                })

            if quantity <= 0 or price <= 0:
                issues.append({
                    "severity": "ERROR",
                    "issue_type": "NON_POSITIVE_QTY_OR_PRICE",
                    "account_id": account_id,
                    "symbol": symbol,
                    "trade_date": trade_date,
                    "row_number": idx,
                    "message": f"Row {idx} has non-positive quantity ({quantity}) or price ({price}).",
                    "raw_payload": str(row),
                })
                continue

            if exec_id:
                unique_key = f"{account_id}|{exec_id}"
            else:
                unique_key = f"{account_id}|{trade_date}|{trade_time}|{symbol}|{norm_side}|{quantity}|{price:.4f}|{order_id}"
                issues.append({
                    "severity": "WARNING",
                    "issue_type": "COMPOSITE_KEY_USED",
                    "account_id": account_id,
                    "symbol": symbol,
                    "trade_date": trade_date,
                    "row_number": idx,
                    "message": f"Row {idx} lacked an execution/reference ID; using composite uniqueness key '{unique_key}'.",
                })

            charges = compute_execution_charges(norm_side, quantity, price, trade_value, settings)

            record = {
                "account_id": account_id,
                "trade_date": trade_date,
                "trade_time": trade_time,
                "terminal_id": terminal_id,
                "market_type": market_type,
                "symbol": symbol,
                "raw_side": raw_side,
                "normalized_side": norm_side,
                "order_id": order_id,
                "execution_id": exec_id,
                "status": status,
                "quantity": quantity,
                "remaining_qty": remaining_qty,
                "order_qty": order_qty,
                "price": price,
                "trade_value": trade_value,
                "commission_slab": charges["commission_slab"],
                "commission_amount": charges["commission_amount"],
                "sst_amount": charges["sst_amount"],
                "total_charges": charges["total_charges"],
                "unique_trade_key": unique_key,
                "source_filename": filename,
                "row_number": idx,
            }

            if unique_key in seen_keys:
                issues.append({
                    "severity": "WARNING",
                    "issue_type": "DUPLICATE_IN_FILE",
                    "account_id": account_id,
                    "symbol": symbol,
                    "trade_date": trade_date,
                    "row_number": idx,
                    "message": f"Duplicate execution key '{unique_key}' inside {filename} at row {idx}; skipped.",
                    "raw_payload": str(row),
                })
                continue
            seen_keys.add(unique_key)

            if account_id not in AUTHORIZED_ACCOUNTS:
                unexpected_executions.append(record)
                issues.append({
                    "severity": "WARNING",
                    "issue_type": "UNEXPECTED_ACCOUNT",
                    "account_id": account_id,
                    "symbol": symbol,
                    "trade_date": trade_date,
                    "row_number": idx,
                    "message": f"Unexpected Account '{account_id}' isolated on row {idx} ({symbol}, {raw_side} {quantity} @ {price}). Excluded from 91404 and 91269 books.",
                    "raw_payload": str(row),
                })
            else:
                executions.append(record)
                detected_dates.add(trade_date)

        except Exception as e:
            issues.append({
                "severity": "ERROR",
                "issue_type": "PARSE_EXCEPTION",
                "row_number": idx,
                "message": f"Failed to parse row {idx}: {e}",
                "raw_payload": str(row),
            })

    return {
        "filename": filename,
        "detected_dates": sorted(detected_dates),
        "rows_read": rows_read,
        "executions": executions,
        "unexpected_executions": unexpected_executions,
        "unexpected_account_rows": unexpected_executions,
        "issues": issues,
    }


def parse_legacy_account_workbook(xlsx_path: Path) -> Dict[str, Any]:
    """
    Parse historical symbol-level daily blocks from 'Daily Trade Summary' in
    Trade_Summary_Account_91404_Updated_*.xlsx or Trade_Summary_Account_91269_Updated_*.xlsx.
    """
    wb = openpyxl.load_workbook(xlsx_path, read_only=False, data_only=True)
    ws_settings = wb["Settings"] if "Settings" in wb.sheetnames else None
    account_id = None
    if "91404" in xlsx_path.name:
        account_id = "91404"
    elif "91269" in xlsx_path.name:
        account_id = "91269"

    ws = wb["Daily Trade Summary"]
    max_r = ws.max_row
    blocks: List[Dict[str, Any]] = []
    symbol_rows: List[Dict[str, Any]] = []

    r = 1
    while r <= max_r:
        v1 = ws.cell(r, 1).value
        if isinstance(v1, str) and v1.startswith("Daily Trade Summary —"):
            raw_date = ws.cell(r + 2, 2).value
            if isinstance(raw_date, datetime):
                trade_date = raw_date.strftime("%Y-%m-%d")
            else:
                trade_date = parse_tradelog_date(str(raw_date))
            acc_cell = str(ws.cell(r + 3, 2).value or "")
            if not account_id:
                m_acc = re.search(r"(91404|91269)", acc_cell)
                if m_acc:
                    account_id = m_acc.group(1)
            market_str = str(ws.cell(r + 4, 2).value or "FUT")
            source_file = str(ws.cell(r + 7, 2).value or xlsx_path.name)

            # Find 'Symbol' header row below r
            sym_hdr_r = r + 10
            while sym_hdr_r <= min(r + 16, max_r):
                if ws.cell(sym_hdr_r, 1).value == "Symbol":
                    break
                sym_hdr_r += 1

            curr_r = sym_hdr_r + 1
            block_syms = []
            while curr_r <= max_r:
                sym_val = ws.cell(curr_r, 1).value
                if sym_val is None or str(sym_val).strip() == "":
                    curr_r += 1
                    continue
                if str(sym_val).strip().upper() == "TOTAL":
                    total_row = {
                        "buy_qty": int(round(float(ws.cell(curr_r, 2).value or 0))),
                        "buy_value": float(ws.cell(curr_r, 3).value or 0.0),
                        "sell_qty": int(round(float(ws.cell(curr_r, 5).value or 0))),
                        "sell_value": float(ws.cell(curr_r, 6).value or 0.0),
                        "gross_pnl": float(ws.cell(curr_r, 10).value or 0.0),
                        "commission_amount": float(ws.cell(curr_r, 12).value or 0.0),
                        "sst_amount": float(ws.cell(curr_r, 13).value or 0.0),
                        "total_charges": float(ws.cell(curr_r, 14).value or 0.0),
                        "net_pnl": float(ws.cell(curr_r, 15).value or 0.0),
                    }
                    blocks.append({
                        "account_id": account_id,
                        "trade_date": trade_date,
                        "source_file": source_file,
                        "symbols_count": len(block_syms),
                        "totals": total_row,
                    })
                    r = curr_r + 1
                    break

                b_qty = int(round(float(ws.cell(curr_r, 2).value or 0)))
                b_val = float(ws.cell(curr_r, 3).value or 0.0)
                avg_b = ws.cell(curr_r, 4).value
                s_qty = int(round(float(ws.cell(curr_r, 5).value or 0)))
                s_val = float(ws.cell(curr_r, 6).value or 0.0)
                avg_s = ws.cell(curr_r, 7).value
                p_diff = ws.cell(curr_r, 8).value
                net_q = int(round(float(ws.cell(curr_r, 9).value or (b_qty - s_qty))))
                gross = float(ws.cell(curr_r, 10).value or (s_val - b_val))
                slab = str(ws.cell(curr_r, 11).value or "")
                comm = float(ws.cell(curr_r, 12).value or 0.0)
                sst = float(ws.cell(curr_r, 13).value or 0.0)
                t_chg = float(ws.cell(curr_r, 14).value or (comm + sst))
                net_p = float(ws.cell(curr_r, 15).value or (gross - t_chg))
                rob = ws.cell(curr_r, 16).value

                row_dict = {
                    "account_id": account_id,
                    "trade_date": trade_date,
                    "symbol": str(sym_val).strip().upper(),
                    "market_type": "FUT" if "-" in str(sym_val) else "Reg",
                    "buy_qty": b_qty,
                    "buy_value": b_val,
                    "avg_buy_price": float(avg_b) if isinstance(avg_b, (int, float)) else (b_val / b_qty if b_qty > 0 else None),
                    "sell_qty": s_qty,
                    "sell_value": s_val,
                    "avg_sell_price": float(avg_s) if isinstance(avg_s, (int, float)) else (s_val / s_qty if s_qty > 0 else None),
                    "price_diff": float(p_diff) if isinstance(p_diff, (int, float)) else None,
                    "net_qty": net_q,
                    "gross_pnl": gross,
                    "commission_slab_label": slab,
                    "commission_amount": comm,
                    "sst_amount": sst,
                    "total_charges": t_chg,
                    "net_pnl": net_p,
                    "return_on_buy": float(rob) if isinstance(rob, (int, float)) else (net_p / b_val if b_val > 0 else None),
                    "origin_type": "LEGACY_MIGRATED",
                    "source_filename": source_file,
                }
                symbol_rows.append(row_dict)
                block_syms.append(row_dict)
                curr_r += 1
            else:
                r = curr_r + 1
        else:
            r += 1

    wb.close()
    return {
        "account_id": account_id,
        "blocks": blocks,
        "symbol_rows": symbol_rows,
    }

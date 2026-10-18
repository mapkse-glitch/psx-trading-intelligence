import csv
import io
import re
from datetime import datetime
from pathlib import Path
from typing import Dict, Any, List, Optional
import openpyxl
from psx_app.config import TICKER_ALIASES


def normalize_ticker(raw_ticker: str) -> str:
    t = raw_ticker.strip().upper()
    return TICKER_ALIASES.get(t, t)


def parse_psx_date(raw_date: str) -> Optional[str]:
    """Convert DD/MM/YYYY, MM/DD/YYYY, or YYYY-MM-DD into ISO YYYY-MM-DD."""
    s = raw_date.strip()
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%d-%b-%Y", "%m/%d/%Y"):
        try:
            return datetime.strptime(s, fmt).strftime("%Y-%m-%d")
        except ValueError:
            continue
    return None


def parse_top5_txt(content: str, filename: str = "unknown.txt") -> Dict[str, Any]:
    """
    Parse a daily ticker Price Summary TXT file (Module A).
    Header example: 'ATRL Price Summary for 16/10/2026 as per 16/10/2026, 16:43:23'
    Row format: '"1,221.17","1","20","24,423.40"' -> (Rate, Trades, Volume, TradedValue)
    """
    clean_content = content.lstrip("\ufeff")
    lines = [ln.strip() for ln in clean_content.splitlines() if ln.strip()]
    issues: List[Dict[str, Any]] = []

    if not lines:
        return {
            "ticker": normalize_ticker(Path(filename).stem),
            "trade_date": None,
            "snapshot_time": None,
            "rows_read": 0,
            "valid_rows": 0,
            "top5": [],
            "issues": [{"severity": "ERROR", "issue_type": "EMPTY_FILE", "message": f"{filename} is empty."}],
        }

    header = lines[0]
    m = re.match(
        r"^([A-Za-z0-9_-]+)\s+Price Summary for\s+(\d{2}/\d{2}/\d{4})(?:\s+as per\s+\d{2}/\d{2}/\d{4},\s*([0-9:]+))?",
        header,
        re.IGNORECASE,
    )
    file_ticker = normalize_ticker(re.sub(r"[_-].*$", "", Path(filename).stem))
    if m:
        ticker = normalize_ticker(m.group(1))
        trade_date = datetime.strptime(m.group(2), "%d/%m/%Y").strftime("%Y-%m-%d")
        snapshot_time = m.group(3)
        data_lines = lines[1:]
        if file_ticker and file_ticker != ticker and not file_ticker.startswith(ticker):
            issues.append({
                "severity": "WARNING",
                "issue_type": "TICKER_FILENAME_MISMATCH",
                "symbol": ticker,
                "trade_date": trade_date,
                "message": f"Header ticker '{ticker}' differs from filename stem '{file_ticker}'; using header ticker '{ticker}'.",
            })
    else:
        ticker = file_ticker
        # Try extracting date from filename
        dm = re.search(r"(\d{2}[-/]\d{2}[-/]\d{4}|\d{4}-\d{2}-\d{2})", filename)
        trade_date = parse_psx_date(dm.group(1)) if dm else None
        snapshot_time = None
        # Check if line 0 is already data
        data_lines = lines if lines[0].startswith('"') or lines[0][0].isdigit() else lines[1:]
        if not trade_date:
            issues.append({
                "severity": "ERROR",
                "issue_type": "MISSING_DATE",
                "symbol": ticker,
                "message": f"Could not determine trading date from header or filename in {filename}.",
            })

    parsed_rows: List[Dict[str, Any]] = []
    reader = csv.reader(io.StringIO("\n".join(data_lines)))
    rows_read = 0
    for idx, row in enumerate(reader, start=2 if m else 1):
        if not row or all(not c.strip() for c in row):
            continue
        rows_read += 1
        if len(row) < 3:
            issues.append({
                "severity": "WARNING",
                "issue_type": "PARSE_FAILURE",
                "symbol": ticker,
                "trade_date": trade_date,
                "row_number": idx,
                "message": f"Row {idx} has fewer than 3 columns: {row}",
                "raw_payload": str(row),
            })
            continue
        try:
            rate = float(row[0].replace(",", "").strip())
            trades = int(float(row[1].replace(",", "").strip()))
            volume = int(float(row[2].replace(",", "").strip()))
            if rate <= 0 or volume <= 0 or trades <= 0:
                issues.append({
                    "severity": "WARNING",
                    "issue_type": "NON_POSITIVE_VALUES",
                    "symbol": ticker,
                    "trade_date": trade_date,
                    "row_number": idx,
                    "message": f"Row {idx} has non-positive rate/volume/trades: {row}",
                    "raw_payload": str(row),
                })
                continue
            parsed_rows.append({
                "rate": rate,
                "volume": volume,
                "trades": trades,
                "row_number": idx,
            })
        except Exception as e:
            issues.append({
                "severity": "WARNING",
                "issue_type": "INVALID_NUMERIC",
                "symbol": ticker,
                "trade_date": trade_date,
                "row_number": idx,
                "message": f"Row {idx} failed numeric validation ({e}): {row}",
                "raw_payload": str(row),
            })

    # Rank by Volume DESC, then Trades DESC, then Rate DESC
    parsed_rows.sort(key=lambda r: (r["volume"], r["trades"], r["rate"]), reverse=True)

    if len(parsed_rows) < 5:
        issues.append({
            "severity": "WARNING",
            "issue_type": "FEWER_THAN_5_ROWS",
            "symbol": ticker,
            "trade_date": trade_date,
            "message": f"{ticker} on {trade_date} has only {len(parsed_rows)} valid rows (< 5); retaining available rows without inventing data.",
        })

    selected = parsed_rows[:5]
    # Check for volume ties within top 5 or at the 5th/6th boundary
    for i in range(len(selected) - 1):
        if selected[i]["volume"] == selected[i + 1]["volume"]:
            issues.append({
                "severity": "INFO",
                "issue_type": "VOLUME_TIE",
                "symbol": ticker,
                "trade_date": trade_date,
                "message": f"Tie in Volume ({selected[i]['volume']}) between rank {i+1} (rate {selected[i]['rate']}) and rank {i+2} (rate {selected[i+1]['rate']}).",
            })
    if len(parsed_rows) > 5 and parsed_rows[4]["volume"] == parsed_rows[5]["volume"]:
        issues.append({
            "severity": "WARNING",
            "issue_type": "BOUNDARY_VOLUME_TIE",
            "symbol": ticker,
            "trade_date": trade_date,
            "message": f"Boundary tie at Rank 5/6 with Volume={parsed_rows[4]['volume']} for {ticker} on {trade_date}.",
        })

    top5_records = []
    for rank_idx, item in enumerate(selected, start=1):
        top5_records.append({
            "trade_date": trade_date,
            "ticker": ticker,
            "rank": rank_idx,
            "rate": item["rate"],
            "volume": item["volume"],
            "trades": item["trades"],
            "snapshot_time": snapshot_time,
            "source_filename": filename,
        })

    mmddyyyy = datetime.strptime(trade_date, "%Y-%m-%d").strftime("%m/%d/%Y") if trade_date else None
    return {
        "ticker": ticker,
        "trade_date": trade_date,
        "mmddyyyy": mmddyyyy,
        "snapshot_time": snapshot_time,
        "rows_read": rows_read,
        "valid_rows": len(parsed_rows),
        "top5": top5_records,
        "issues": issues,
    }


def parse_legacy_top5_workbook(xlsx_path: Path) -> Dict[str, Any]:
    """Parse existing PSX_Top5_Volume_Rate_*.xlsx workbook for historical migration."""
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    records: List[Dict[str, Any]] = []
    for sheet in wb.sheetnames:
        ticker = normalize_ticker(sheet)
        ws = wb[sheet]
        current_date: Optional[str] = None
        for row in ws.iter_rows(values_only=True):
            if not row or all(x is None for x in row):
                continue
            first = row[0]
            if isinstance(first, str) and "Top 5 Volume Trades" in first:
                m = re.search(r"\((\d{2}/\d{2}/\d{4})\)", first)
                if m:
                    # Legacy workbook uses MM/DD/YYYY in block header
                    current_date = datetime.strptime(m.group(1), "%m/%d/%Y").strftime("%Y-%m-%d")
                continue
            if first == "Rank":
                continue
            if current_date and isinstance(first, (int, float)) and 1 <= int(first) <= 5:
                records.append({
                    "trade_date": current_date,
                    "ticker": ticker,
                    "rank": int(first),
                    "rate": float(row[1]),
                    "volume": int(round(float(row[2]))),
                    "trades": int(round(float(row[3]))),
                    "snapshot_time": None,
                    "source_filename": xlsx_path.name,
                    "origin_type": "LEGACY_WORKBOOK",
                })
    wb.close()
    return {"records": records, "sheets_count": len(wb.sheetnames)}

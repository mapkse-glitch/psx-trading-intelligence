import hashlib,json,re,io
from pathlib import Path
import openpyxl
from psx_app.database.schema import get_connection

BUNDLED = Path(__file__).resolve().parents[1]/'reference'/'sectors.json'

def parse_reference_workbook(raw):
    if len(raw)>2_000_000: raise ValueError('Reference workbooks must be smaller than 2 MB.')
    w=openpyxl.load_workbook(io.BytesIO(raw),read_only=True,data_only=True)
    if sum((s.max_row or 0)*(s.max_column or 0) for s in w)>250000: raise ValueError('Reference workbook exceeds the safe cell limit.')
    lists={}; warnings=[]
    for s in w:
        rows=s.iter_rows(values_only=True);header=next(rows,())
        names={str(v or '').strip().upper():i for i,v in enumerate(header) if v}
        if not {'SYMBOL','COMPANY NAME','SECTOR'}<=names.keys():
            warnings.append(s.title+': unsupported headers; left as custom data, not imported.');continue
        items={}
        for row in rows:
            if any(v is not None for i,v in enumerate(row) if i not in [names[k] for k in ('SYMBOL','COMPANY NAME','SECTOR')]):
                if s.title+': unlabelled/extra columns ignored; no automatic aliases.' not in warnings: warnings.append(s.title+': unlabelled/extra columns ignored; no automatic aliases.')
            vals=[str(row[names[k]] or '').strip() if names[k]<len(row) else '' for k in ('SYMBOL','COMPANY NAME','SECTOR')]
            if not vals[0]: continue
            sym=vals[0].upper()
            if not re.fullmatch(r'[A-Z0-9._-]{1,24}',sym) or not vals[1] or not vals[2]: raise ValueError('Invalid company row in '+s.title)
            item=dict(symbol=sym,company=vals[1][:250],sector=vals[2][:150])
            if sym in items and items[sym]!=item: raise ValueError('Conflicting duplicate symbol '+sym)
            items[sym]=item
        if items: lists[s.title]=list(items.values())
    w.close()
    if not lists: raise ValueError('Custom workbook: map columns to SYMBOL, COMPANY NAME, SECTOR and upload again. No price data changed.')
    if len(lists.get('KSE100',[]))!=100 and 'KSE100' in lists: warnings.append('KSE100 list is not exactly 100 symbols. Membership is user-supplied, not independently verified.')
    if 'COMP_LIST' not in lists: warnings.append('No COMP_LIST sheet: company directory will remain unchanged.')
    return {"lists": lists, "warnings": warnings}

class SectorReference:
    def __init__(self, db_path=None): self.db_path=db_path
    def lists(self):
        base=json.loads(BUNDLED.read_text())
        with get_connection(self.db_path) as c:
            rows=c.execute('SELECT list_name,symbol,company,sector FROM company_reference ORDER BY list_name,symbol').fetchall()
        for name in {r['list_name'] for r in rows}: base[name]=[]
        for r in rows: base[r['list_name']].append({k:r[k] for k in ('symbol','company','sector')})
        return base
    def directory(self): return {r['symbol']:r for r in self.lists().get('COMP_LIST',[])}
    def preview(self, raw):
        parsed = parse_reference_workbook(raw)
        lists, warnings = parsed['lists'], parsed['warnings']
        with get_connection(self.db_path) as c: known={r[0] for r in c.execute('SELECT DISTINCT symbol FROM equity_ohlcv')}
        old=self.lists()
        changes={name:sum(1 for r in values if r not in old.get(name,[])) for name,values in lists.items()}
        token=hashlib.sha256(raw).hexdigest()
        payload={'lists':lists,'warnings':warnings,'changes':changes,'unmatched_symbols':sorted({r['symbol'] for values in lists.values() for r in values}-known)}
        with get_connection(self.db_path) as c:
            c.execute('INSERT OR REPLACE INTO reference_previews(token,payload_json) VALUES (?,?)',(token,json.dumps(payload)))
            c.execute('DELETE FROM reference_previews WHERE token NOT IN (SELECT token FROM reference_previews ORDER BY rowid DESC LIMIT 20)')
            c.commit()
        return {'token':token,'counts':{k:len(v) for k,v in lists.items()},**{k:v for k,v in payload.items() if k!='lists'}}
    def commit(self, token):
        with get_connection(self.db_path) as c:
            row=c.execute('SELECT payload_json FROM reference_previews WHERE token=?',(token,)).fetchone()
            if not row: raise ValueError('Preview expired or missing; preview the workbook again.')
            payload=json.loads(row[0])
            for name,items in payload['lists'].items():
                c.execute('DELETE FROM company_reference WHERE list_name=?',(name,))
                c.executemany('INSERT INTO company_reference(list_name,symbol,company,sector) VALUES (?,?,?,?)',[(name,r['symbol'],r['company'],r['sector']) for r in items])
            c.execute('DELETE FROM reference_previews WHERE token=?',(token,))
            c.execute("INSERT INTO audit_events(event_type,actor,details_json,created_at) VALUES ('SECTOR_REFERENCE_IMPORTED','dashboard',?,datetime('now'))",(json.dumps({'token':token,'counts':{k:len(v) for k,v in payload['lists'].items()}}),))
            c.commit()
        return {'status':'COMMITTED','warnings':payload['warnings']}

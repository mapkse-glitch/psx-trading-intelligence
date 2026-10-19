import csv
import io
import re
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional
import openpyxl
from openpyxl.formatting.rule import CellIsRule, DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from psx_app.config import CORE_36_TICKERS, DATA_ROOT, DB_PATH, ensure_directories
from psx_app.database.schema import get_connection
from psx_app.services.account_service import AccountService
from psx_app.services.futures_service import FuturesService
from psx_app.services.reconciliation_service import ReconciliationService
from psx_app.services.technical_service import TechnicalService

NAVY_HEADER_FILL = PatternFill("solid", fgColor="1B365D")
SUBHEADER_FILL = PatternFill("solid", fgColor="2C5282")
TOTAL_FILL = PatternFill("solid", fgColor="EDF2F7")
INPUT_FILL = PatternFill("solid", fgColor="FFFBEB")
ZEBRA_FILL = PatternFill("solid", fgColor="F8FAFC")
GREEN_FILL = PatternFill("solid", fgColor="DEF7EC")
RED_FILL = PatternFill("solid", fgColor="FDE8E8")

WHITE_BOLD = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
TITLE_FONT = Font(name="Calibri", size=13, bold=True, color="1B365D")
SUBTITLE_FONT = Font(name="Calibri", size=9, italic=True, color="4A5568")
BOLD_FONT = Font(name="Calibri", size=10, bold=True, color="1A202C")
REG_FONT = Font(name="Calibri", size=10, bold=False, color="1A202C")
GREEN_FONT = Font(name="Calibri", size=10, bold=True, color="046C4E")
RED_FONT = Font(name="Calibri", size=10, bold=True, color="C81E1E")

THIN_BORDER = Border(
    left=Side(style="thin", color="CBD5E0"),
    right=Side(style="thin", color="CBD5E0"),
    top=Side(style="thin", color="CBD5E0"),
    bottom=Side(style="thin", color="CBD5E0"),
)
TOTAL_BORDER = Border(
    top=Side(style="thin", color="1B365D"),
    bottom=Side(style="double", color="1B365D"),
    left=Side(style="thin", color="CBD5E0"),
    right=Side(style="thin", color="CBD5E0"),
)


def _auto_fit_columns(ws: Any, min_w: int = 11, max_w: int = 34) -> None:
    for col_cells in ws.columns:
        col_letter = get_column_letter(col_cells[0].column)
        max_len = 0
        for cell in col_cells[:250]:
            v = cell.value
            if v is not None:
                s = str(v)
                if not s.startswith("="):
                    max_len = max(max_len, len(s))
        ws.column_dimensions[col_letter].width = max(min_w, min(max_w, max_len + 4))


class ExportService:
    """
    Layer 8 — Legacy-Compatible Excel (.xlsx) and CSV Export Generator.
    """

    def __init__(self, db_path: Optional[Path] = None, data_root: Optional[Path] = None):
        self.db_path = db_path or DB_PATH
        self.data_root = data_root or DATA_ROOT
        ensure_directories(self.data_root)
        self.exports_dir = self.data_root / "exports"
        self.account_service = AccountService(self.db_path)
        self.futures_service = FuturesService(self.db_path)
        self.technical_service = TechnicalService(self.db_path)
        self.reconciliation_service = ReconciliationService(self.db_path)

    # ------------------------------------------------------------------
    # MODULE A: Top 5 Volume & Rate Register Workbook
    # ------------------------------------------------------------------
    def export_top5_workbook(self) -> Path:
        wb = openpyxl.Workbook()
        wb.remove(wb.active)

        with get_connection(self.db_path) as c:
            tickers = [
                r[0]
                for r in c.execute("SELECT DISTINCT ticker FROM top5_volume_records ORDER BY ticker ASC").fetchall()
            ]
            all_rows = [
                dict(r)
                for r in c.execute(
                    """
                    SELECT trade_date, ticker, rank, rate, volume, trades
                    FROM top5_volume_records
                    ORDER BY ticker ASC, trade_date ASC, rank ASC
                    """
                ).fetchall()
            ]

        by_ticker_date: Dict[str, Dict[str, List[Dict[str, Any]]]] = {}
        for r in all_rows:
            by_ticker_date.setdefault(r["ticker"], {}).setdefault(r["trade_date"], []).append(r)

        if not tickers:
            ws = wb.create_sheet("NO_DATA")
            ws["A1"] = "No Top 5 records available."
        else:
            for tkr in tickers:
                ws = wb.create_sheet(title=tkr[:31])
                date_map = by_ticker_date.get(tkr, {})
                curr_row = 1
                block_idx = 0

                for iso_date in sorted(date_map.keys()):
                    block_idx += 1
                    dt_str = datetime.strptime(iso_date, "%Y-%m-%d").strftime("%m/%d/%Y")
                    title_cell = ws.cell(curr_row, 1, f"{tkr} — Top 5 Volume Trades ({dt_str})")
                    title_cell.font = TITLE_FONT
                    curr_row += 2

                    hdr_row = curr_row
                    headers = ["Rank", "Rate", "Volume", "Trades"]
                    for col_i, h in enumerate(headers, start=1):
                        c = ws.cell(hdr_row, col_i, h)
                        c.font = WHITE_BOLD
                        c.fill = NAVY_HEADER_FILL
                        c.alignment = Alignment(horizontal="center")

                    curr_row += 1
                    data_start = curr_row
                    for item in date_map[iso_date]:
                        c_rk = ws.cell(curr_row, 1, int(item["rank"]))
                        c_rt = ws.cell(curr_row, 2, float(item["rate"]))
                        c_vl = ws.cell(curr_row, 3, int(item["volume"]))
                        c_tr = ws.cell(curr_row, 4, int(item["trades"]))

                        c_rk.number_format = "0"
                        c_rt.number_format = "#,##0.00"
                        c_vl.number_format = "#,##0"
                        c_tr.number_format = "#,##0"

                        c_rk.alignment = Alignment(horizontal="center")
                        for cell in (c_rk, c_rt, c_vl, c_tr):
                            cell.font = REG_FONT
                            cell.border = THIN_BORDER
                        curr_row += 1
                    data_end = curr_row - 1

                    if data_end >= data_start:
                        tbl_name = f"T_{re.sub(r'[^A-Za-z0-9]', '', tkr)}_Block{block_idx}"
                        ref = f"A{hdr_row}:D{data_end}"
                        tab = Table(displayName=tbl_name, ref=ref)
                        tab.tableStyleInfo = TableStyleInfo(
                            name="TableStyleMedium2",
                            showFirstColumn=False,
                            showLastColumn=False,
                            showRowStripes=True,
                            showColumnStripes=False,
                        )
                        ws.add_table(tab)
                        # Conditional data bar on Volume
                        ws.conditional_formatting.add(
                            f"C{data_start}:C{data_end}",
                            DataBarRule(start_type="min", end_type="max", color="3182CE"),
                        )

                    curr_row += 1  # Blank spacer row between date blocks

                ws.column_dimensions["A"].width = 12
                ws.column_dimensions["B"].width = 16
                ws.column_dimensions["C"].width = 18
                ws.column_dimensions["D"].width = 14

        out_path = self.exports_dir / "PSX_Top5_Volume_Rate_Register.xlsx"
        wb.save(out_path)
        return out_path

    # ------------------------------------------------------------------
    # MODULE B: Account-Specific Workbook (91404 or 91269)
    # ------------------------------------------------------------------
    def export_account_workbook(self, account_id: str, include_audit_tabs: bool = True) -> Path:
        rep = self.account_service.get_account_report(account_id)
        wb = openpyxl.Workbook()
        wb.remove(wb.active)

        dash = rep["dashboard"]
        latest_date = dash["latest_trading_day"] or "2026-10-16"

        # 1. Dashboard Sheet
        ws_d = wb.create_sheet("Dashboard")
        ws_d["A1"] = "Trading Performance Dashboard"
        ws_d["A1"].font = TITLE_FONT
        ws_d["A2"] = f"Account {account_id} (THM00601) • Market FUT, Reg • Currency PKR • Generated {datetime.now().strftime('%d-%b-%Y %H:%M')}"
        ws_d["A2"].font = SUBTITLE_FONT
        ws_d["A4"] = "Latest Trading Day:"
        ws_d["A4"].font = BOLD_FONT
        ws_d["B4"] = latest_date
        ws_d["B4"].font = BOLD_FONT

        ws_d["A6"] = "LATEST TRADING DAY"
        ws_d["A6"].font = WHITE_BOLD
        ws_d["A6"].fill = NAVY_HEADER_FILL
        ws_d["B6"].fill = NAVY_HEADER_FILL

        ld = dash["latest_day"] or {}
        metrics_left = [
            ("Buy Value", ld.get("buy_value", 0.0), "#,##0.00"),
            ("Sell Value", ld.get("sell_value", 0.0), "#,##0.00"),
            ("Gross P&L", ld.get("gross_pnl", 0.0), "#,##0.00"),
            ("Commission", ld.get("commission_amount", 0.0), "#,##0.00"),
            ("SST (13%)", ld.get("sst_amount", 0.0), "#,##0.00"),
            ("Total Charges", ld.get("total_charges", 0.0), "#,##0.00"),
            ("Net P&L", ld.get("net_pnl", 0.0), "#,##0.00"),
            ("Return on Buy", ld.get("return_on_buy") or 0.0, "0.00%"),
        ]
        for idx, (m_lbl, m_val, fmt) in enumerate(metrics_left, start=7):
            c1 = ws_d.cell(idx, 1, m_lbl)
            c2 = ws_d.cell(idx, 2, m_val)
            c1.font = BOLD_FONT
            c2.font = REG_FONT
            c2.number_format = fmt
            c1.border = THIN_BORDER
            c2.border = THIN_BORDER

        ws_d["D6"] = "PERFORMANCE SNAPSHOT"
        ws_d["D6"].font = WHITE_BOLD
        ws_d["D6"].fill = NAVY_HEADER_FILL
        ws_d["E6"].fill = NAVY_HEADER_FILL
        ws_d["F6"].fill = NAVY_HEADER_FILL

        for col_i, h in enumerate(["Period", "Net P&L (PKR)", "Return on Buy"], start=4):
            c = ws_d.cell(7, col_i, h)
            c.font = BOLD_FONT
            c.fill = TOTAL_FILL
            c.border = THIN_BORDER

        periods_right = [
            ("MTD", dash["mtd"]),
            ("QTD", dash["qtd"]),
            ("YTD", dash["ytd"]),
        ]
        for idx, (p_lbl, p_obj) in enumerate(periods_right, start=8):
            c_p = ws_d.cell(idx, 4, p_lbl)
            c_n = ws_d.cell(idx, 5, (p_obj or {}).get("net_pnl", 0.0))
            c_r = ws_d.cell(idx, 6, (p_obj or {}).get("return_on_buy") or 0.0)
            c_n.number_format = "#,##0.00"
            c_r.number_format = "0.00%"
            for cell in (c_p, c_n, c_r):
                cell.font = REG_FONT
                cell.border = THIN_BORDER

        ws_d["D12"] = "ACTIVITY — ALL IMPORTED DATA"
        ws_d["D12"].font = WHITE_BOLD
        ws_d["D12"].fill = NAVY_HEADER_FILL
        ws_d["E12"].fill = NAVY_HEADER_FILL

        act_rows = [
            ("Total Traded Volume (PKR)", dash["total_traded_turnover_pkr"], "#,##0.00"),
            ("Total Traded Shares", dash["total_traded_shares"], "#,##0"),
            ("Symbols", dash["symbols_count"], "#,##0"),
            ("Trading Days", dash["trading_days_count"], "#,##0"),
        ]
        for idx, (a_lbl, a_val, fmt) in enumerate(act_rows, start=13):
            c1 = ws_d.cell(idx, 4, a_lbl)
            c2 = ws_d.cell(idx, 5, a_val)
            c1.font = BOLD_FONT
            c2.font = REG_FONT
            c2.number_format = fmt
            c1.border = THIN_BORDER
            c2.border = THIN_BORDER

        ws_d["A18"] = "BEST / WORST SYMBOL (Net P&L, all dates)"
        ws_d["A18"].font = WHITE_BOLD
        ws_d["A18"].fill = NAVY_HEADER_FILL
        if dash["best_symbol"]:
            ws_d["A19"] = "Best Symbol"
            ws_d["B19"] = dash["best_symbol"]["symbol"]
            ws_d["C19"] = dash["best_symbol"]["net_pnl"]
            ws_d["C19"].number_format = "#,##0.00"
            ws_d["C19"].font = GREEN_FONT
        if dash["worst_symbol"]:
            ws_d["A20"] = "Worst Symbol"
            ws_d["B20"] = dash["worst_symbol"]["symbol"]
            ws_d["C20"] = dash["worst_symbol"]["net_pnl"]
            ws_d["C20"].number_format = "#,##0.00"
            ws_d["C20"].font = RED_FONT
        _auto_fit_columns(ws_d)

        # 2. Daily Summary Sheet
        ws_ds = wb.create_sheet("Daily Summary")
        ws_ds["A1"] = "Daily Summary"
        ws_ds["A1"].font = TITLE_FONT
        ws_ds["A2"] = "One row per trading date — aggregated from underlying transactions."
        ws_ds["A2"].font = SUBTITLE_FONT
        ds_headers = [
            "Date", "Buy Qty", "Sell Qty", "Net Qty", "Buy Value (PKR)", "Sell Value (PKR)",
            "Avg Buy Price", "Avg Sell Price", "Gross P&L (PKR)", "Commission (PKR)",
            "SST (PKR)", "Total Charges (PKR)", "Net P&L (PKR)", "Return on Buy"
        ]
        for col_i, h in enumerate(ds_headers, start=1):
            c = ws_ds.cell(4, col_i, h)
            c.font = WHITE_BOLD
            c.fill = NAVY_HEADER_FILL
        r_idx = 5
        for row in rep["daily_summary"]["rows"]:
            vals = [
                row["trade_date"], row["buy_qty"], row["sell_qty"], row["net_qty"],
                row["buy_value"], row["sell_value"],
                row["avg_buy_price"] if row["avg_buy_price"] is not None else "N/A",
                row["avg_sell_price"] if row["avg_sell_price"] is not None else "N/A",
                row["gross_pnl"], row["commission_amount"], row["sst_amount"],
                row["total_charges"], row["net_pnl"],
                row["return_on_buy"] if row["return_on_buy"] is not None else "N/A",
            ]
            for c_i, v in enumerate(vals, start=1):
                cell = ws_ds.cell(r_idx, c_i, v)
                cell.font = REG_FONT
                cell.border = THIN_BORDER
                if c_i in (2, 3, 4):
                    cell.number_format = "#,##0"
                elif c_i in (5, 6, 7, 8, 9, 10, 11, 12, 13) and isinstance(v, (int, float)):
                    cell.number_format = "#,##0.00"
                elif c_i == 14 and isinstance(v, (int, float)):
                    cell.number_format = "0.00%"
            r_idx += 1
        if rep["daily_summary"]["total"]:
            tot = rep["daily_summary"]["total"]
            t_vals = [
                "TOTAL", tot["buy_qty"], tot["sell_qty"], tot["net_qty"],
                tot["buy_value"], tot["sell_value"],
                tot["avg_buy_price"] if tot["avg_buy_price"] is not None else "N/A",
                tot["avg_sell_price"] if tot["avg_sell_price"] is not None else "N/A",
                tot["gross_pnl"], tot["commission_amount"], tot["sst_amount"],
                tot["total_charges"], tot["net_pnl"],
                tot["return_on_buy"] if tot["return_on_buy"] is not None else "N/A",
            ]
            for c_i, v in enumerate(t_vals, start=1):
                cell = ws_ds.cell(r_idx, c_i, v)
                cell.font = BOLD_FONT
                cell.fill = TOTAL_FILL
                cell.border = TOTAL_BORDER
                if c_i in (2, 3, 4):
                    cell.number_format = "#,##0"
                elif c_i in (5, 6, 7, 8, 9, 10, 11, 12, 13) and isinstance(v, (int, float)):
                    cell.number_format = "#,##0.00"
                elif c_i == 14 and isinstance(v, (int, float)):
                    cell.number_format = "0.00%"
        ws_ds.freeze_panes = "A5"
        _auto_fit_columns(ws_ds)

        # 3, 4, 5. MTD, QTD, YTD Summary Sheets
        for tab_title, obj_key in [
            ("MTD Summary", "mtd_summary"),
            ("QTD Summary", "qtd_summary"),
            ("YTD Summary", "ytd_summary"),
        ]:
            ws_p = wb.create_sheet(tab_title)
            p_data = rep[obj_key]
            ws_p["A1"] = f"{tab_title} — Account {account_id}"
            ws_p["A1"].font = TITLE_FONT
            ws_p["A4"] = "Period Start (calendar):"
            ws_p["B4"] = p_data.get("start")
            ws_p["A5"] = "Period End (latest trade):"
            ws_p["B5"] = p_data.get("end")
            ws_p["A6"] = "First Trading Day in Period:"
            ws_p["B6"] = p_data.get("first_trading_day")
            ws_p["A7"] = "Trading Days:"
            ws_p["B7"] = p_data.get("trading_days", 0)

            ws_p["A9"] = "Metric"
            ws_p["B9"] = "Value"
            ws_p["A9"].font = WHITE_BOLD
            ws_p["B9"].font = WHITE_BOLD
            ws_p["A9"].fill = NAVY_HEADER_FILL
            ws_p["B9"].fill = NAVY_HEADER_FILL

            kv_rows = [
                ("Buy Quantity", p_data.get("buy_qty", 0), "#,##0"),
                ("Sell Quantity", p_data.get("sell_qty", 0), "#,##0"),
                ("Net Quantity", p_data.get("net_qty", 0), "#,##0"),
                ("Buy Value (PKR)", p_data.get("buy_value", 0.0), "#,##0.00"),
                ("Average Buy Price", p_data.get("avg_buy_price") or "N/A", "#,##0.00"),
                ("Sell Value (PKR)", p_data.get("sell_value", 0.0), "#,##0.00"),
                ("Average Sell Price", p_data.get("avg_sell_price") or "N/A", "#,##0.00"),
                ("Gross P&L (PKR)", p_data.get("gross_pnl", 0.0), "#,##0.00"),
                ("Commission (PKR)", p_data.get("commission_amount", 0.0), "#,##0.00"),
                ("SST (PKR)", p_data.get("sst_amount", 0.0), "#,##0.00"),
                ("Total Charges (PKR)", p_data.get("total_charges", 0.0), "#,##0.00"),
                ("Net P&L (PKR)", p_data.get("net_pnl", 0.0), "#,##0.00"),
                ("Return on Buy", p_data.get("return_on_buy") if p_data.get("return_on_buy") is not None else "N/A", "0.00%"),
            ]
            for idx_k, (k_lbl, k_val, fmt) in enumerate(kv_rows, start=10):
                c1 = ws_p.cell(idx_k, 1, k_lbl)
                c2 = ws_p.cell(idx_k, 2, k_val)
                c1.font = BOLD_FONT
                c2.font = REG_FONT
                c1.border = THIN_BORDER
                c2.border = THIN_BORDER
                if isinstance(k_val, (int, float)):
                    c2.number_format = fmt
            _auto_fit_columns(ws_p)

        # 6 & 7. Monthly Summary & Yearly Summary
        for tab_title, obj_key, col1_name, key_field in [
            ("Monthly Summary", "monthly_summary", "Month", "month"),
            ("Yearly Summary", "yearly_summary", "Year", "year"),
        ]:
            ws_my = wb.create_sheet(tab_title)
            ws_my["A1"] = f"{tab_title} — Account {account_id}"
            ws_my["A1"].font = TITLE_FONT
            ws_my["A2"] = "Accumulated from complete underlying transaction history (weighted averages; Return on Buy = Net P&L / Buy Value)."
            ws_my["A2"].font = SUBTITLE_FONT
            headers_16 = [
                col1_name, "Buy Qty", "Buy Value (PKR)", "Avg Buy Price",
                "Sell Qty", "Sell Value (PKR)", "Avg Sell Price",
                "Price Diff (Sell−Buy)", "Net Qty", "Gross P&L (PKR)",
                "Commission Rate (Slab)", "Commission (PKR)", "SST 13% (PKR)",
                "Total Charges (PKR)", "Net P&L (PKR)", "Return on Buy"
            ]
            for c_i, h in enumerate(headers_16, start=1):
                c = ws_my.cell(4, c_i, h)
                c.font = WHITE_BOLD
                c.fill = NAVY_HEADER_FILL
            r_i = 5
            for row in rep[obj_key]["rows"]:
                vals = [
                    row[key_field], row["buy_qty"], row["buy_value"],
                    row["avg_buy_price"] if row["avg_buy_price"] is not None else "N/A",
                    row["sell_qty"], row["sell_value"],
                    row["avg_sell_price"] if row["avg_sell_price"] is not None else "N/A",
                    row["price_diff"] if row["price_diff"] is not None else "N/A",
                    row["net_qty"], row["gross_pnl"], row["commission_slab_label"],
                    row["commission_amount"], row["sst_amount"], row["total_charges"],
                    row["net_pnl"], row["return_on_buy"] if row["return_on_buy"] is not None else "N/A",
                ]
                for c_i, v in enumerate(vals, start=1):
                    cell = ws_my.cell(r_i, c_i, v)
                    cell.font = REG_FONT
                    cell.border = THIN_BORDER
                    if c_i in (2, 5, 9):
                        cell.number_format = "#,##0"
                    elif c_i in (3, 4, 6, 7, 8, 10, 12, 13, 14, 15) and isinstance(v, (int, float)):
                        cell.number_format = "#,##0.00"
                    elif c_i == 16 and isinstance(v, (int, float)):
                        cell.number_format = "0.00%"
                r_i += 1
            if rep[obj_key]["total"]:
                tot = rep[obj_key]["total"]
                t_vals = [
                    "TOTAL", tot["buy_qty"], tot["buy_value"],
                    tot["avg_buy_price"] if tot["avg_buy_price"] is not None else "N/A",
                    tot["sell_qty"], tot["sell_value"],
                    tot["avg_sell_price"] if tot["avg_sell_price"] is not None else "N/A",
                    tot["price_diff"] if tot["price_diff"] is not None else "N/A",
                    tot["net_qty"], tot["gross_pnl"], None,
                    tot["commission_amount"], tot["sst_amount"], tot["total_charges"],
                    tot["net_pnl"], tot["return_on_buy"] if tot["return_on_buy"] is not None else "N/A",
                ]
                for c_i, v in enumerate(t_vals, start=1):
                    cell = ws_my.cell(r_i, c_i, v)
                    cell.font = BOLD_FONT
                    cell.fill = TOTAL_FILL
                    cell.border = TOTAL_BORDER
                    if c_i in (2, 5, 9):
                        cell.number_format = "#,##0"
                    elif c_i in (3, 4, 6, 7, 8, 10, 12, 13, 14, 15) and isinstance(v, (int, float)):
                        cell.number_format = "#,##0.00"
                    elif c_i == 16 and isinstance(v, (int, float)):
                        cell.number_format = "0.00%"
            _auto_fit_columns(ws_my)

        # 8. Symbol Performance Sheet
        ws_sp = wb.create_sheet("Symbol Performance")
        ws_sp["A1"] = "Symbol Performance — All Trading Dates"
        ws_sp["A1"].font = TITLE_FONT
        sp_headers = [
            "Symbol", "Buy Qty", "Sell Qty", "Buy Value (PKR)", "Sell Value (PKR)",
            "Gross P&L (PKR)", "Commission (PKR)", "SST (PKR)", "Net P&L (PKR)", "Return on Buy", "Note"
        ]
        for c_i, h in enumerate(sp_headers, start=1):
            c = ws_sp.cell(4, c_i, h)
            c.font = WHITE_BOLD
            c.fill = NAVY_HEADER_FILL
        for idx_s, row in enumerate(rep["symbol_performance"], start=5):
            note = "◄ Best performer" if idx_s == 5 else ("◄ Worst performer" if idx_s == 4 + len(rep["symbol_performance"]) else None)
            vals = [
                row["symbol"], row["buy_qty"], row["sell_qty"], row["buy_value"], row["sell_value"],
                row["gross_pnl"], row["commission_amount"], row["sst_amount"], row["net_pnl"],
                row["return_on_buy"] if row["return_on_buy"] is not None else "N/A", note
            ]
            for c_i, v in enumerate(vals, start=1):
                cell = ws_sp.cell(idx_s, c_i, v)
                cell.font = REG_FONT
                cell.border = THIN_BORDER
                if c_i in (2, 3):
                    cell.number_format = "#,##0"
                elif c_i in (4, 5, 6, 7, 8, 9) and isinstance(v, (int, float)):
                    cell.number_format = "#,##0.00"
                elif c_i == 10 and isinstance(v, (int, float)):
                    cell.number_format = "0.00%"
        _auto_fit_columns(ws_sp)

        # 9. Daily Trade Summary Sheet (stacked date blocks matching legacy layout)
        ws_dts = wb.create_sheet("Daily Trade Summary")
        curr_r = 1
        sym_headers_16 = [
            "Symbol", "Buy Qty", "Buy Value (PKR)", "Avg Buy Price",
            "Sell Qty", "Sell Value (PKR)", "Avg Sell Price",
            "Price Diff (Sell−Buy)", "Net Qty", "Gross P&L (PKR)",
            "Commission Rate (Slab)", "Commission (PKR)", "SST 13% (PKR)",
            "Total Charges (PKR)", "Net P&L (PKR)", "Return on Buy"
        ]
        for d_iso, blk in rep["all_daily_blocks"].items():
            dt_fmt = datetime.strptime(d_iso, "%Y-%m-%d").strftime("%d %b %Y")
            b_sum = blk["summary"]
            b_syms = blk["symbols"]

            ws_dts.cell(curr_r, 1, f"Daily Trade Summary — {dt_fmt}").font = TITLE_FONT
            ws_dts.cell(curr_r + 1, 1, "Gross intraday buy/sell flow • commission & SST • net P&L (all values PKR)").font = SUBTITLE_FONT
            ws_dts.cell(curr_r + 1, 4, "Day Results").font = BOLD_FONT

            ws_dts.cell(curr_r + 2, 1, "Trading Date:")
            ws_dts.cell(curr_r + 2, 2, d_iso)
            ws_dts.cell(curr_r + 3, 1, "Account:")
            ws_dts.cell(curr_r + 3, 2, f"THM00601 (A/c {account_id})")
            ws_dts.cell(curr_r + 4, 1, "Currency:")
            ws_dts.cell(curr_r + 4, 2, "PKR")
            ws_dts.cell(curr_r + 5, 1, "Source File:")
            ws_dts.cell(curr_r + 5, 2, b_syms[0]["source_filename"] if b_syms else "TradeLog")

            day_kpis = [
                ("Total Buy Value", b_sum["buy_value"]),
                ("Total Sell Value", b_sum["sell_value"]),
                ("Gross P&L", b_sum["gross_pnl"]),
                ("Commission", b_sum["commission_amount"]),
                ("SST (13%)", b_sum["sst_amount"]),
                ("Total Charges", b_sum["total_charges"]),
                ("Net P&L", b_sum["net_pnl"]),
                ("Return on Buy", b_sum["return_on_buy"] or 0.0),
                ("Net Position (shares)", b_sum["net_qty"]),
            ]
            for k_idx, (k_lbl, k_v) in enumerate(day_kpis):
                ws_dts.cell(curr_r + 2 + k_idx, 4, k_lbl).font = BOLD_FONT
                c_kv = ws_dts.cell(curr_r + 2 + k_idx, 5, k_v)
                c_kv.number_format = "0.00%" if k_lbl == "Return on Buy" else ("#,##0" if "Position" in k_lbl else "#,##0.00")

            hdr_r = curr_r + 12
            for c_i, h in enumerate(sym_headers_16, start=1):
                c = ws_dts.cell(hdr_r, c_i, h)
                c.font = WHITE_BOLD
                c.fill = NAVY_HEADER_FILL

            curr_r = hdr_r + 1
            for s_row in b_syms:
                vals = [
                    s_row["symbol"], s_row["buy_qty"], s_row["buy_value"],
                    s_row["avg_buy_price"] if s_row["avg_buy_price"] is not None else "N/A",
                    s_row["sell_qty"], s_row["sell_value"],
                    s_row["avg_sell_price"] if s_row["avg_sell_price"] is not None else "N/A",
                    s_row["price_diff"] if s_row["price_diff"] is not None else "N/A",
                    s_row["net_qty"], s_row["gross_pnl"], s_row["commission_slab_label"],
                    s_row["commission_amount"], s_row["sst_amount"], s_row["total_charges"],
                    s_row["net_pnl"], s_row["return_on_buy"] if s_row["return_on_buy"] is not None else "N/A",
                ]
                for c_i, v in enumerate(vals, start=1):
                    cell = ws_dts.cell(curr_r, c_i, v)
                    cell.font = REG_FONT
                    cell.border = THIN_BORDER
                    if c_i in (2, 5, 9):
                        cell.number_format = "#,##0"
                    elif c_i in (3, 4, 6, 7, 8, 10, 12, 13, 14, 15) and isinstance(v, (int, float)):
                        cell.number_format = "#,##0.00"
                    elif c_i == 16 and isinstance(v, (int, float)):
                        cell.number_format = "0.00%"
                curr_r += 1

            t_vals = [
                "TOTAL", b_sum["buy_qty"], b_sum["buy_value"],
                b_sum["avg_buy_price"] if b_sum["avg_buy_price"] is not None else "N/A",
                b_sum["sell_qty"], b_sum["sell_value"],
                b_sum["avg_sell_price"] if b_sum["avg_sell_price"] is not None else "N/A",
                b_sum["price_diff"] if b_sum["price_diff"] is not None else "N/A",
                b_sum["net_qty"], b_sum["gross_pnl"], None,
                b_sum["commission_amount"], b_sum["sst_amount"], b_sum["total_charges"],
                b_sum["net_pnl"], b_sum["return_on_buy"] if b_sum["return_on_buy"] is not None else "N/A",
            ]
            for c_i, v in enumerate(t_vals, start=1):
                cell = ws_dts.cell(curr_r, c_i, v)
                cell.font = BOLD_FONT
                cell.fill = TOTAL_FILL
                cell.border = TOTAL_BORDER
                if c_i in (2, 5, 9):
                    cell.number_format = "#,##0"
                elif c_i in (3, 4, 6, 7, 8, 10, 12, 13, 14, 15) and isinstance(v, (int, float)):
                    cell.number_format = "#,##0.00"
                elif c_i == 16 and isinstance(v, (int, float)):
                    cell.number_format = "0.00%"
            curr_r += 3
        _auto_fit_columns(ws_dts)

        # 10. Settings Sheet
        ws_st = wb.create_sheet("Settings")
        ws_st["A1"] = "Settings & Assumptions"
        ws_st["A1"].font = TITLE_FONT
        for c_i, h in enumerate(["Setting", "Value", "Notes"], start=1):
            c = ws_st.cell(4, c_i, h)
            c.font = WHITE_BOLD
            c.fill = NAVY_HEADER_FILL
        st_rows = [
            ("Account Number", account_id, f"Specific account for this workbook ({account_id})"),
            ("Commission Slab 1 Upper Rate", rep["settings"]["commission_slab1_upper"], "Sell rate < 54.00 PKR -> Slab 1"),
            ("Commission Slab 1 Rate", rep["settings"]["commission_slab1_paisa_per_share"], "0.0175 paisa per share -> PKR = Sell Qty * 0.0175 / 100"),
            ("Commission Slab 2 Starting Rate", rep["settings"]["commission_slab2_start"], "Sell rate >= 54.00 PKR -> Slab 2"),
            ("Commission Slab 2 Rate", rep["settings"]["commission_slab2_rate"], "0.0325% of sell value (0.000325)"),
            ("SST Rate", rep["settings"]["sst_rate"], "13% of commission amount"),
            ("Currency", rep["settings"]["currency"], "Pakistan Rupee (PKR)"),
        ]
        for idx_r, (s_k, s_v, s_n) in enumerate(st_rows, start=5):
            ws_st.cell(idx_r, 1, s_k).font = BOLD_FONT
            ws_st.cell(idx_r, 2, s_v).font = REG_FONT
            ws_st.cell(idx_r, 3, s_n).font = REG_FONT
        _auto_fit_columns(ws_st)

        if include_audit_tabs:
            # 11. Trade Log Sheet
            ws_tl = wb.create_sheet("Trade Log")
            ws_tl["A1"] = f"Master Trade Log — Account {account_id}"
            ws_tl["A1"].font = TITLE_FONT
            tl_headers = [
                "Trading Date", "Time", "Account", "Symbol", "Market",
                "Original Type", "Normalized Side", "Order ID", "Execution ID",
                "Quantity", "Price (PKR)", "Trade Value (PKR)", "Commission Slab",
                "Commission (PKR)", "SST 13% (PKR)", "Total Charges (PKR)",
                "Unique Trade Key", "Source Filename"
            ]
            for c_i, h in enumerate(tl_headers, start=1):
                c = ws_tl.cell(3, c_i, h)
                c.font = WHITE_BOLD
                c.fill = NAVY_HEADER_FILL
            for idx_e, ex in enumerate(rep["trade_log"], start=4):
                vals = [
                    ex["trade_date"], ex["trade_time"], ex["account_id"], ex["symbol"],
                    ex["market_type"], ex["raw_side"], ex["normalized_side"], ex["order_id"],
                    ex["execution_id"], ex["quantity"], ex["price"], ex["trade_value"],
                    ex["commission_slab"], ex["commission_amount"], ex["sst_amount"],
                    ex["total_charges"], ex["unique_trade_key"], ex["source_filename"],
                ]
                for c_i, v in enumerate(vals, start=1):
                    cell = ws_tl.cell(idx_e, c_i, v)
                    cell.font = REG_FONT
                    cell.border = THIN_BORDER
            _auto_fit_columns(ws_tl)

            # 12. Data Checks Sheet
            ws_dc = wb.create_sheet("Data Checks")
            ws_dc["A1"] = f"Data Checks & Reconciliation — Account {account_id}"
            ws_dc["A1"].font = TITLE_FONT
            rec = rep["data_checks"]["reconciliation"]
            ws_dc["A3"] = "Reconciliation Passed:"
            ws_dc["B3"] = "YES" if rec["passed"] else "NO"
            ws_dc["A4"] = "Daily vs Symbol Net P&L Diff:"
            ws_dc["B4"] = rec["daily_vs_symbol_net_pnl_diff"]
            ws_dc["A5"] = "Monthly vs Daily Net P&L Diff:"
            ws_dc["B5"] = rec["monthly_vs_daily_net_pnl_diff"]
            ws_dc["A6"] = "Yearly vs Daily Net P&L Diff:"
            ws_dc["B6"] = rec["yearly_vs_daily_net_pnl_diff"]
            ws_dc["A7"] = "Unmatched Net Qty Symbols Count:"
            ws_dc["B7"] = rec["unmatched_count"]
            _auto_fit_columns(ws_dc)

        out_path = self.exports_dir / f"Trade_Summary_Account_{account_id}.xlsx"
        wb.save(out_path)
        return out_path

    # ------------------------------------------------------------------
    # MODULE C: Futures Comparison & Annualized Analysis Workbook
    # Strictly 7 visible worksheets + 1 hidden 'Data' worksheet
    # ------------------------------------------------------------------
    def export_futures_workbook(
        self,
        selected_date: Optional[str] = None,
        selected_contract: Optional[str] = None,
        selected_symbol: Optional[str] = None,
        selected_future_item: Optional[str] = None,
        days: int = 30,
    ) -> Path:
        meta = self.futures_service.get_metadata(apply_rolling=True)
        sel_date = selected_date or meta["latest_date"]
        sel_contract = selected_contract or ("OCT" if "OCT" in meta["contracts"] else (meta["contracts"][0] if meta["contracts"] else "OCT"))
        sel_symbol = selected_symbol or ("OGDC" if "OGDC" in meta["base_symbols"] else (meta["base_symbols"][0] if meta["base_symbols"] else "OGDC"))
        sel_item = selected_future_item or ("AGHA-OCT" if "AGHA-OCT" in meta["contract_symbols"] else (meta["contract_symbols"][0] if meta["contract_symbols"] else "OGDC-OCT"))

        wb = openpyxl.Workbook()
        wb.remove(wb.active)

        comp_headers = [
            "Date", "Symbol", "Contract/Month", "Current Value", "Future Value",
            "Volume", "High", "Low", "Difference", "Difference %", "Days", "Annualized %"
        ]

        def populate_comparison_sheet(
            ws: Any,
            title_text: str,
            rows_data: List[Dict[str, Any]],
            first_filter_label: str = "Selected Date:",
            first_filter_val: Optional[str] = None,
            extra_filter_label: Optional[str] = None,
            extra_filter_val: Optional[str] = None,
        ) -> None:
            ws["A2"] = title_text
            ws["A2"].font = TITLE_FONT
            ws["B4"] = first_filter_label
            ws["B4"].font = BOLD_FONT
            ws["C4"] = first_filter_val if first_filter_val is not None else sel_date
            ws["C4"].fill = INPUT_FILL
            ws["C4"].font = BOLD_FONT
            ws["E4"] = "Days:"
            ws["E4"].font = BOLD_FONT
            ws["F4"] = int(days)
            ws["F4"].fill = INPUT_FILL
            ws["F4"].font = BOLD_FONT
            if extra_filter_label:
                ws["H4"] = extra_filter_label
                ws["H4"].font = BOLD_FONT
                ws["I4"] = extra_filter_val
                ws["I4"].fill = INPUT_FILL
                ws["I4"].font = BOLD_FONT

            for c_i, h in enumerate(comp_headers, start=1):
                c = ws.cell(6, c_i, h)
                c.font = WHITE_BOLD
                c.fill = NAVY_HEADER_FILL
                c.alignment = Alignment(horizontal="center")

            for idx_r, r in enumerate(rows_data, start=7):
                diff_pct_excel = (r["difference_pct"] / 100.0) if isinstance(r["difference_pct"], (int, float)) else "N/A"
                ann_pct_excel = (r["annualized_pct"] / 100.0) if isinstance(r["annualized_pct"], (int, float)) else "N/A"
                vals = [
                    r["trade_date"], r["symbol"], r["contract"],
                    r["current_value"], r["future_value"], r["volume"],
                    r["high"], r["low"], r["difference"],
                    diff_pct_excel, r["days"], ann_pct_excel,
                ]
                for c_i, v in enumerate(vals, start=1):
                    cell = ws.cell(idx_r, c_i, v)
                    cell.font = REG_FONT
                    cell.border = THIN_BORDER
                    if c_i in (4, 5, 7, 8, 9) and isinstance(v, (int, float)):
                        cell.number_format = "#,##0.00"
                    elif c_i in (6, 11) and isinstance(v, (int, float)):
                        cell.number_format = "#,##0"
                    elif c_i in (10, 12) and isinstance(v, (int, float)):
                        cell.number_format = "0.00%"
                    if c_i in (9, 10, 12) and isinstance(v, (int, float)):
                        if v > 0:
                            cell.font = GREEN_FONT
                        elif v < 0:
                            cell.font = RED_FONT
            ws.freeze_panes = "A7"
            _auto_fit_columns(ws)

        # Sheets 1..5
        sf_contract = "OCTB" if "OCTB" in meta["contracts"] else sel_contract
        mc_data = self.futures_service.get_comparison_rows("monthly_comparison", selected_date=sel_date, days=days)["rows"]
        mf_data = self.futures_service.get_comparison_rows("month_filter", selected_date=sel_date, contract_month=sel_contract, days=days)["rows"]
        sf_data = self.futures_service.get_comparison_rows("symbol_filter", selected_date="ALL", contract_month=sf_contract, base_symbol=sel_symbol, days=days)["rows"]
        ac_data = self.futures_service.get_comparison_rows("annualized_comparison", selected_date=sel_date, days=days)["rows"]
        df_data = self.futures_service.get_comparison_rows("date_filter", selected_date=sel_date, days=days)["rows"]

        populate_comparison_sheet(wb.create_sheet("Monthly Comparison"), "MONTHLY COMPARISON — FUTURES vs CURRENT", mc_data)
        populate_comparison_sheet(wb.create_sheet("Month Filter"), "MONTH / CONTRACT FILTER", mf_data, "Selected Date:", sel_date, "Month/Contract:", sel_contract)
        populate_comparison_sheet(wb.create_sheet("Symbol Filter"), "SYMBOL FILTER — ALL DATES FOR SELECTED CONTRACT & SYMBOL", sf_data, "Month/Contract:", sf_contract, "Symbol:", sel_symbol)
        populate_comparison_sheet(wb.create_sheet("Annualized Comparison"), "ANNUALIZED COMPARISON — ALL SYMBOLS & CONTRACTS", ac_data)
        populate_comparison_sheet(wb.create_sheet("Date Filter"), "DATE FILTER — FULL DAY COMPARISON", df_data)

        # Sheet 6: Future Symbol Summary
        ws_fss = wb.create_sheet("Future Symbol Summary")
        fss_res = self.futures_service.get_future_symbol_summary(contract_symbol=sel_item)
        ws_fss["A2"] = "FUTURE SYMBOL SUMMARY — ALL RETAINED DATES (2026F)"
        ws_fss["A2"].font = TITLE_FONT
        ws_fss["B4"] = "Symbol / Item:"
        ws_fss["B4"].font = BOLD_FONT
        ws_fss["C4"] = fss_res["selected_contract_symbol"]
        ws_fss["C4"].fill = INPUT_FILL
        ws_fss["C4"].font = BOLD_FONT

        fss_headers = ["DATE", "PR_CL", "OPEN", "High", "LOW", "AVG", "CLOSE", "VOLUME", "Change"]
        for c_i, h in enumerate(fss_headers, start=1):
            c = ws_fss.cell(6, c_i, h)
            c.font = WHITE_BOLD
            c.fill = NAVY_HEADER_FILL
        for idx_r, r in enumerate(fss_res["rows"], start=7):
            vals = [r["DATE"], r["PR_CL"], r["OPEN"], r["High"], r["LOW"], r["AVG"], r["CLOSE"], r["VOLUME"], r["Change"]]
            for c_i, v in enumerate(vals, start=1):
                cell = ws_fss.cell(idx_r, c_i, v)
                cell.font = REG_FONT
                cell.border = THIN_BORDER
                if c_i in (2, 3, 4, 5, 6, 7, 9) and isinstance(v, (int, float)):
                    cell.number_format = "#,##0.00"
                elif c_i == 8 and isinstance(v, (int, float)):
                    cell.number_format = "#,##0"
                if c_i == 9 and isinstance(v, (int, float)):
                    cell.font = GREEN_FONT if v > 0 else (RED_FONT if v < 0 else REG_FONT)
        ws_fss.freeze_panes = "A7"
        _auto_fit_columns(ws_fss)

        # Sheet 7: Date Summary
        ws_ds = wb.create_sheet("Date Summary")
        ds_res = self.futures_service.get_date_summary(selected_date=sel_date)
        ws_ds["A2"] = "DATE SUMMARY — ALL SYMBOLS FOR SELECTED DATE (2026F)"
        ws_ds["A2"].font = TITLE_FONT
        ws_ds["B4"] = "Selected Date:"
        ws_ds["B4"].font = BOLD_FONT
        ws_ds["C4"] = sel_date
        ws_ds["C4"].fill = INPUT_FILL
        ws_ds["C4"].font = BOLD_FONT

        ds_headers = ["DATE", "SYMBOL", "PR_CL", "OPEN", "High", "LOW", "AVG", "CLOSE", "VOLUME", "Change"]
        for c_i, h in enumerate(ds_headers, start=1):
            c = ws_ds.cell(6, c_i, h)
            c.font = WHITE_BOLD
            c.fill = NAVY_HEADER_FILL
        if not ds_res["rows"]:
            ws_ds.cell(7, 1, "No data available for this date.").font = RED_FONT
        else:
            for idx_r, r in enumerate(ds_res["rows"], start=7):
                vals = [r["DATE"], r["SYMBOL"], r["PR_CL"], r["OPEN"], r["High"], r["LOW"], r["AVG"], r["CLOSE"], r["VOLUME"], r["Change"]]
                for c_i, v in enumerate(vals, start=1):
                    cell = ws_ds.cell(idx_r, c_i, v)
                    cell.font = REG_FONT
                    cell.border = THIN_BORDER
                    if c_i in (3, 4, 5, 6, 7, 8, 10) and isinstance(v, (int, float)):
                        cell.number_format = "#,##0.00"
                    elif c_i == 9 and isinstance(v, (int, float)):
                        cell.number_format = "#,##0"
                    if c_i == 10 and isinstance(v, (int, float)):
                        cell.font = GREEN_FONT if v > 0 else (RED_FONT if v < 0 else REG_FONT)
        ws_ds.freeze_panes = "A7"
        _auto_fit_columns(ws_ds)

        # Sheet 8: Hidden 'Data' sheet (strictly 1 hidden worksheet)
        ws_hidden = wb.create_sheet("Data")
        ws_hidden.sheet_state = "hidden"
        ws_hidden.append([
            "Date", "Symbol", "Contract", "Contract_Symbol",
            "Current_Value", "Future_Value", "Volume", "High", "Low", "PR_CL", "Open", "Avg"
        ])
        cutoff = meta["cutoff_date"]
        with get_connection(self.db_path) as c:
            q_clause = "WHERE f.trade_date >= ?" if cutoff else ""
            q_params = (cutoff,) if cutoff else ()
            h_rows = c.execute(
                f"""
                SELECT f.trade_date, f.base_symbol, f.contract_month, f.contract_symbol,
                       e.close AS current_val, f.close AS future_val, f.volume, f.high, f.low,
                       f.pr_close, f.open, f.avg_price
                FROM futures_ohlcv f
                LEFT JOIN equity_ohlcv e
                    ON e.trade_date = f.trade_date AND e.symbol = f.base_symbol
                {q_clause}
                ORDER BY f.trade_date ASC, f.contract_symbol ASC
                """,
                q_params,
            ).fetchall()
            for hr in h_rows:
                ws_hidden.append(list(hr))

        out_path = self.exports_dir / "Futures_Analysis.xlsx"
        wb.save(out_path)
        return out_path

    # ------------------------------------------------------------------
    # MODULE D: Technical Analysis & 36-Ticker Workbook
    # ------------------------------------------------------------------
    def export_technical_workbook(self) -> Path:
        wb = openpyxl.Workbook()
        wb.remove(wb.active)

        uni = self.technical_service.get_universe_summary(scope="all")
        ws_sum = wb.create_sheet("SUMMARY")
        ws_sum["A1"] = f"PSX TECHNICAL SUMMARY — ALL {uni['count']} TICKERS | RSI(14), Daily & Weekly Pivots, MACD(12,26,9), Bollinger(20,2)"
        ws_sum["A1"].font = TITLE_FONT
        ws_sum["A2"] = "Generated from canonical equity_ohlcv database | Daily & Weekly Pivots | Explainable 9-Factor Signal Engine"
        ws_sum["A2"].font = SUBTITLE_FONT

        sum_headers = [
            "SYMBOL", "LAST CLOSE", "CHG %", "RSI(14)", "ZONE", "TREND vs SMA20",
            "PIVOT", "R1", "R2", "R3", "S1", "S2", "S3",
            "W-PIVOT", "W-R1", "W-R2", "W-R3", "W-S1", "W-S2", "W-S3",
            "MACD", "MACD SIG", "MACD HIST", "BB UPPER", "BB LOWER", "%B",
            "COMBINED SIGNAL", "CHART TAB"
        ]
        for c_i, h in enumerate(sum_headers, start=1):
            c = ws_sum.cell(6, c_i, h)
            c.font = WHITE_BOLD
            c.fill = NAVY_HEADER_FILL

        for idx_r, item in enumerate(uni["items"], start=7):
            dp = item["daily_pivots"]["classic"]
            wp = (item["weekly_pivots"] or {}).get("pivots") or {}
            vals = [
                item["symbol"], item["close"], item["change_pct"] / 100.0,
                item["rsi14"] if item["rsi14"] is not None else "NO DATA",
                item["rsi_zone"], item["trend_vs_sma20"],
                dp["P"], dp["R1"], dp["R2"], dp["R3"], dp["S1"], dp["S2"], dp["S3"],
                wp.get("P"), wp.get("R1"), wp.get("R2"), wp.get("R3"), wp.get("S1"), wp.get("S2"), wp.get("S3"),
                item["macd"], item["macd_sig"], item["macd_hist"],
                item["bb_up"], item["bb_low"], item["pct_b"],
                item["signal"]["combined_signal"],
                item["symbol"] if item["symbol"] in CORE_36_TICKERS else "—",
            ]
            for c_i, v in enumerate(vals, start=1):
                cell = ws_sum.cell(idx_r, c_i, v)
                cell.font = REG_FONT
                cell.border = THIN_BORDER
                if c_i == 3 and isinstance(v, (int, float)):
                    cell.number_format = "0.00%"
                elif isinstance(v, float):
                    cell.number_format = "#,##0.00"
        ws_sum.freeze_panes = "A7"
        _auto_fit_columns(ws_sum)

        ws_tl = wb.create_sheet("TICKER_LIST")
        ws_tl.sheet_state = "hidden"
        for item in uni["items"]:
            ws_tl.append([item["symbol"]])

        # Build the 36 core ticker tabs + hidden _DATA tabs
        for sym in CORE_36_TICKERS:
            t_res = self.technical_service.get_ticker_analysis(sym, tail_bars=250)
            latest = t_res["latest"]
            if not latest:
                continue
            ws_t = wb.create_sheet(sym)
            ws_t["A1"] = f"{sym} — PSX TECHNICAL SHEET | RSI(14), PIVOTS & 9-INDICATOR DASHBOARD"
            ws_t["A1"].font = TITLE_FONT
            ws_t["A2"] = f"Data: {latest['first_date']} -> {latest['latest_date']} ({latest['bars_count']} sessions) | Last Close: {latest['close']:.2f} PKR"
            ws_t["A2"].font = SUBTITLE_FONT

            ws_t["A4"] = "PIVOT POINTS"
            ws_t["A4"].font = BOLD_FONT
            for c_i, h in enumerate(["LEVEL", "CLASSIC", "FIBONACCI", "CAMARILLA"], start=1):
                c = ws_t.cell(5, c_i, h)
                c.font = WHITE_BOLD
                c.fill = NAVY_HEADER_FILL

            dp = latest["daily_pivots"]
            for idx_lvl, lvl in enumerate(["R3", "R2", "R1", "P", "S1", "S2", "S3"], start=6):
                ws_t.cell(idx_lvl, 1, "PIVOT (P)" if lvl == "P" else lvl).font = BOLD_FONT
                ws_t.cell(idx_lvl, 2, dp["classic"][lvl]).number_format = "#,##0.00"
                ws_t.cell(idx_lvl, 3, dp["fibonacci"][lvl]).number_format = "#,##0.00"
                ws_t.cell(idx_lvl, 4, dp["camarilla"][lvl]).number_format = "#,##0.00"

            ws_t["A15"] = "INDICATOR DASHBOARD — Latest Value / Signal / Advice"
            ws_t["A15"].font = BOLD_FONT
            for c_i, h in enumerate(["INDICATOR", "VALUE", "SIGNAL", "ADVICE", "VOTE"], start=1):
                c = ws_t.cell(16, c_i, h)
                c.font = WHITE_BOLD
                c.fill = NAVY_HEADER_FILL
            for idx_b, b_item in enumerate(latest["signal"]["indicator_breakdown"], start=17):
                ws_t.cell(idx_b, 1, b_item["indicator"]).font = BOLD_FONT
                ws_t.cell(idx_b, 2, b_item["value"])
                ws_t.cell(idx_b, 3, b_item["signal"])
                ws_t.cell(idx_b, 4, b_item["advice"])
                ws_t.cell(idx_b, 5, b_item["vote"])

            ws_t["A27"] = "COMPOSITE SIGNAL"
            ws_t["A27"].font = BOLD_FONT
            ws_t["B27"] = f"{latest['signal']['score']:+d} / 9  ->  {latest['signal']['combined_signal']}"
            ws_t["B27"].font = BOLD_FONT
            ws_t["A28"] = "ENTRY / STOP / TARGET REFERENCE"
            ws_t["B28"] = (
                f"Entry Ref: {latest['signal']['entry_reference']} | "
                f"Stop-Loss Ref: {latest['signal']['stop_loss_reference']} | "
                f"Target Ref: {latest['signal']['target_reference']} | "
                f"R/R: {latest['signal']['risk_reward']}"
            )
            _auto_fit_columns(ws_t)

            ws_td = wb.create_sheet(f"{sym}_DATA")
            ws_td.sheet_state = "hidden"
            ws_td.append([
                "DATE", "OPEN", "HIGH", "LOW", "CLOSE", "VOLUME",
                "RSI14", "MACD", "MACD_SIG", "MACD_HIST",
                "BB_UP", "BB_MID", "BB_LOW", "PCT_B", "BB_WIDTH",
                "EMA9", "EMA21", "EMA50", "EMA200_STRICT", "EMA200_RECURSIVE",
                "VWAP_TYPICAL_DAILY", "VWAP_ROLLING20", "ADX", "PLUS_DI", "MINUS_DI", "SCORE", "ACTION"
            ])
            for s_row in t_res["series"]:
                ws_td.append([
                    s_row["trade_date"], s_row["open"], s_row["high"], s_row["low"], s_row["close"], s_row["volume"],
                    s_row["rsi14"], s_row["macd"], s_row["macd_sig"], s_row["macd_hist"],
                    s_row["bb_up"], s_row["bb_mid"], s_row["bb_low"], s_row["pct_b"], s_row["bb_width"],
                    s_row["ema9"], s_row["ema21"], s_row["ema50"], s_row["ema200"], s_row["ema200_recursive"],
                    s_row["vwap_typical_daily"], s_row["vwap_rolling20"], s_row["adx14"],
                    s_row["plus_di"], s_row["minus_di"], s_row["score"], s_row["action"]
                ])

        out_path = self.exports_dir / "PSX_RSI_Pivot_Analysis.xlsx"
        wb.save(out_path)
        return out_path

    def export_csv_extract(self, module: str) -> Path:
        out_path = self.exports_dir / f"psx_{module.lower()}_extract.csv"
        with get_connection(self.db_path) as c, open(out_path, "w", newline="", encoding="utf-8") as f:
            writer = csv.writer(f)
            if module == "top5":
                rows = c.execute("SELECT trade_date, ticker, rank, rate, volume, trades, source_filename FROM top5_volume_records ORDER BY trade_date DESC, ticker ASC, rank ASC").fetchall()
                writer.writerow(["trade_date", "ticker", "rank", "rate", "volume", "trades", "source_filename"])
                for r in rows:
                    writer.writerow(list(r))
            elif module in ("91404", "91269"):
                rows = c.execute(
                    "SELECT account_id, trade_date, symbol, buy_qty, buy_value, sell_qty, sell_value, gross_pnl, commission_amount, sst_amount, total_charges, net_pnl, return_on_buy FROM account_daily_symbol_summary WHERE account_id=? ORDER BY trade_date DESC, symbol ASC",
                    (module,),
                ).fetchall()
                writer.writerow(["account_id", "trade_date", "symbol", "buy_qty", "buy_value", "sell_qty", "sell_value", "gross_pnl", "commission_amount", "sst_amount", "total_charges", "net_pnl", "return_on_buy"])
                for r in rows:
                    writer.writerow(list(r))
            elif module == "futures":
                rows = c.execute("SELECT trade_date, contract_symbol, base_symbol, contract_month, pr_close, open, high, low, avg_price, close, volume, change_val FROM futures_ohlcv ORDER BY trade_date DESC, contract_symbol ASC").fetchall()
                writer.writerow(["trade_date", "contract_symbol", "base_symbol", "contract_month", "pr_close", "open", "high", "low", "avg_price", "close", "volume", "change_val"])
                for r in rows:
                    writer.writerow(list(r))
            else:
                rows = c.execute("SELECT trade_date, symbol, open, high, low, close, pr_close, volume FROM equity_ohlcv ORDER BY trade_date DESC, symbol ASC LIMIT 20000").fetchall()
                writer.writerow(["trade_date", "symbol", "open", "high", "low", "close", "pr_close", "volume"])
                for r in rows:
                    writer.writerow(list(r))
        return out_path
